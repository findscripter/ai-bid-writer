"""上传与文档解析（FileService）的测试。

全部使用现场生成的合成文档，不读取真实招标材料，也不发起任何网络请求。
"""
from __future__ import annotations

import asyncio
import dataclasses
import io
import os
import random
import tempfile
import threading
import time
import zipfile
from typing import Optional, Sequence

import docx
import fitz
import pytest
from starlette.datastructures import Headers, UploadFile

from app.config import settings
from app.services import file_service
from app.services.file_service import ExtractedDocument, FileService
from app.utils.errors import AppError

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_TYPE = "application/pdf"
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

PARAGRAPH = "本项目为智慧园区综合管理平台建设，投标人须具备相应的软件开发与系统集成能力。"
TABLE_ROWS = [
    ("评分项", "分值", "评分标准"),
    ("技术方案", "30", "方案完整、针对性强"),
    ("商务报价", "40", "按评标基准价计算"),
]
PDF_LINE = "the bidder shall provide a complete technical proposal and service plan."


# ---------------------------------------------------------------------------
# 合成文档
# ---------------------------------------------------------------------------


def _png_bytes(index: int = 0) -> bytes:
    """第 index 张测试图片；每张灰度不同，免得 python-docx 把相同的图片合并成一份。"""
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 40, 40))
    pixmap.clear_with(100 + index * 40)
    return pixmap.tobytes("png")


def make_docx(
    paragraphs: int = 8,
    table: bool = True,
    images: int = 0,
    heading: bool = True,
    table_rows: Sequence[Sequence[str]] = TABLE_ROWS,
) -> bytes:
    """用 python-docx 生成含段落、表格（可选图片）的 Word 文档。"""
    document = docx.Document()
    if heading:
        document.add_heading("第一章 招标公告", level=1)
    for index in range(paragraphs):
        document.add_paragraph(f"第{index + 1}段：{PARAGRAPH}")
    if table:
        document.add_paragraph("表格前的段落。")
        grid = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                grid.cell(row_index, column_index).text = value
        document.add_paragraph("表格后的段落。")
    for index in range(images):
        document.add_picture(io.BytesIO(_png_bytes(index)))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(pages: Sequence[str], lines: int = 20, stamp: Optional[str] = None) -> bytes:
    """用 PyMuPDF 按页生成 PDF。

    'text' 是文字页，'image' 是整页一张图片（模拟扫描页），'blank' 是空白页；
    stamp 会印在每个图片页上，模拟扫描件里零星的页码、水印文字。
    """
    document = fitz.open()
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 200, 280))
    pixmap.clear_with(235)
    for number, kind in enumerate(pages, 1):
        page = document.new_page()
        if kind == "text":
            for line in range(lines):
                page.insert_text(
                    (60, 60 + line * 18),
                    f"Page {number} line {line + 1}: {PDF_LINE}",
                    fontsize=10,
                )
            page.insert_text(
                (60, 80 + lines * 18),
                f"第{number}页 招标文件正文内容示例",
                fontsize=10,
                fontname="china-s",
            )
        elif kind == "image":
            page.insert_image(page.rect, pixmap=pixmap)
            if stamp:
                page.insert_text((500, 800), stamp, fontsize=9)
    data = document.tobytes()
    document.close()
    return data


def make_pdf_with_table() -> bytes:
    """带网格线表格的 PDF，用来验证表格提取。"""
    document = fitz.open()
    page = document.new_page()
    page.insert_text((60, 60), "Scoring table for the technical proposal", fontsize=11)
    left, top, width, height = 60, 100, 150, 30
    for row in range(4):
        page.draw_line((left, top + row * height), (left + 3 * width, top + row * height))
    for column in range(4):
        page.draw_line((left + column * width, top), (left + column * width, top + 3 * height))
    for row in range(3):
        for column in range(3):
            page.insert_text(
                (left + column * width + 8, top + row * height + 20),
                f"R{row + 1}C{column + 1}",
                fontsize=10,
            )
    data = document.tobytes()
    document.close()
    return data


def make_encrypted_pdf(user_password: str) -> bytes:
    document = fitz.open()
    page = document.new_page()
    for line in range(20):
        page.insert_text((60, 60 + line * 18), f"line {line + 1}: {PDF_LINE}", fontsize=10)
    data = document.tobytes(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw=user_password,
        owner_pw="owner-secret",
        permissions=fitz.PDF_PERM_PRINT,
    )
    document.close()
    return data


def make_zip(members: dict) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        for name, data in members.items():
            package.writestr(name, data)
    return buffer.getvalue()


def make_upload(
    data: bytes,
    filename: Optional[str],
    content_type: str = "application/octet-stream",
    size: Optional[int] = None,
) -> UploadFile:
    return UploadFile(
        io.BytesIO(data),
        filename=filename,
        size=size,
        headers=Headers({"content-type": content_type}),
    )


async def extract(data: bytes, filename: Optional[str], **kwargs) -> ExtractedDocument:
    return await FileService.extract_document(make_upload(data, filename, **kwargs))


async def extract_error(data: bytes, filename: Optional[str], **kwargs) -> AppError:
    """抽取必须以使用者可读的 400 错误结束，返回该错误。"""
    with pytest.raises(AppError) as caught:
        await extract(data, filename, **kwargs)
    error = caught.value
    assert error.status_code == 400
    # 解析库的英文报错不能漏给使用者
    for leaked in ("zip", "Traceback", "Error", "Exception", "No /Root"):
        assert leaked not in error.message
    return error


@pytest.fixture(autouse=True)
def no_image_upload(monkeypatch):
    """图片外传默认关闭；开发机上即使配了 IMAGE_UPLOAD_URL 也不能影响测试。"""
    monkeypatch.setattr(FileService, "IMAGE_UPLOAD_URL", "")


@pytest.fixture()
def temp_dir(tmp_path, monkeypatch):
    """把系统临时目录指到测试专用目录，便于检查临时文件有没有清理干净。"""
    directory = tmp_path / "system-temp"
    directory.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(directory))
    return directory


# ---------------------------------------------------------------------------
# 正常抽取
# ---------------------------------------------------------------------------


def test_extracted_document_contract():
    # 其他模块直接依赖这四个字段，名字和顺序都不能变
    names = [item.name for item in dataclasses.fields(ExtractedDocument)]
    assert names == ["filename", "text", "char_count", "warnings"]

    document = ExtractedDocument(filename="a.docx", text="正文", char_count=2)
    assert document.warnings == []


async def test_docx_paragraphs_and_table_are_extracted():
    document = await extract(make_docx(), "招标文件.docx", content_type=DOCX_TYPE)

    assert isinstance(document, ExtractedDocument)
    assert document.filename == "招标文件.docx"
    assert document.char_count == len(document.text) > 0
    assert document.warnings == []

    lines = document.text.splitlines()
    assert "第一章 招标公告" in lines
    # 每个段落独占一行（旧实现会把全部段落用“ | ”拼成一行）
    for index in range(8):
        assert f"第{index + 1}段：{PARAGRAPH}" in lines
    # 表格一行一条记录，单元格内容都在
    for row in TABLE_ROWS:
        assert " | ".join(row) in lines
    assert document.text.count("[表格 1]") == 1
    assert document.text.count("[表格结束]") == 1


async def test_docx_table_stays_at_its_position_in_the_document():
    text = (await extract(make_docx(), "招标文件.docx")).text

    assert (
        text.index("表格前的段落。")
        < text.index("[表格 1]")
        < text.index("评分项 | 分值 | 评分标准")
        < text.index("[表格结束]")
        < text.index("表格后的段落。")
    )


async def test_docx_multi_paragraph_cell_stays_on_one_row():
    source = docx.Document()
    grid = source.add_table(rows=1, cols=2)
    grid.cell(0, 0).text = "评分标准"
    cell = grid.cell(0, 1)
    cell.text = "方案完整得5分"
    cell.add_paragraph("方案较完整得3分")
    for index in range(6):
        source.add_paragraph(f"第{index + 1}段：{PARAGRAPH}")
    buffer = io.BytesIO()
    source.save(buffer)

    document = await extract(buffer.getvalue(), "评分表.docx")

    assert "评分标准 | 方案完整得5分 方案较完整得3分" in document.text.splitlines()


async def test_text_pdf_is_extracted():
    document = await extract(make_pdf(["text"] * 3), "招标文件.pdf", content_type=PDF_TYPE)

    assert document.filename == "招标文件.pdf"
    assert document.char_count == len(document.text) > 0
    assert document.warnings == []
    for number in (1, 2, 3):
        assert f"--- 第 {number} 页 ---" in document.text
        assert f"Page {number} line 20: {PDF_LINE}" in document.text
        assert f"第{number}页 招标文件正文内容示例" in document.text


async def test_pdf_table_is_extracted():
    document = await extract(make_pdf_with_table(), "评分表.pdf")

    lines = document.text.splitlines()
    assert "[表格 1]" in lines
    for row in (1, 2, 3):
        assert f"R{row}C1 | R{row}C2 | R{row}C3" in lines
    assert "[表格结束]" in lines


@pytest.mark.parametrize("filename", ["招标文件.DOCX", "招标文件.Docx"])
async def test_extension_check_is_case_insensitive(filename):
    document = await extract(make_docx(), filename)

    assert f"第1段：{PARAGRAPH}" in document.text


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("C:\\fakepath\\招标文件.docx", "招标文件.docx"),
        ("../../etc/招标文件.docx", "招标文件.docx"),
        ("招标\n文件\u202e.docx", "招标文件.docx"),
        ("招标文件\u3000终版.docx", "招标文件\u3000终版.docx"),
    ],
)
async def test_filename_is_reduced_to_a_display_name(raw, expected, temp_dir):
    document = await extract(make_docx(), raw)

    assert document.filename == expected
    assert os.listdir(temp_dir) == []


async def test_content_type_is_not_trusted():
    # 没装 Office 的机器上浏览器给 .docx 的 Content-Type 可能是空的或者乱的
    document = await extract(make_docx(), "招标文件.docx", content_type="text/plain")
    assert f"第1段：{PARAGRAPH}" in document.text

    # 反过来，Content-Type 写得再对，内容不是 docx 也不行
    await extract_error("只是一段文字".encode("utf-8"), "招标文件.docx", content_type=DOCX_TYPE)


async def test_file_without_extension_is_judged_by_content():
    document = await extract(make_pdf(["text"]), "upload")
    assert f"Page 1 line 1: {PDF_LINE}" in document.text

    document = await extract(make_docx(), None)
    assert document.filename == "未命名文件"
    assert f"第1段：{PARAGRAPH}" in document.text

    error = await extract_error(b"plain text without any signature", "upload")
    assert ".docx" in error.message and ".pdf" in error.message


# ---------------------------------------------------------------------------
# 提示：扫描件、空文档、内容过少
# ---------------------------------------------------------------------------


async def test_scanned_pdf_gets_a_warning_instead_of_silent_success():
    document = await extract(make_pdf(["image"] * 5), "盖章扫描件.pdf", content_type=PDF_TYPE)

    assert document.text == ""  # 不返回只有页码标记的空壳
    assert document.char_count == 0
    assert len(document.warnings) == 1
    warning = document.warnings[0]
    assert "扫描件" in warning
    assert "共 5 页" in warning
    assert "0 个字符" in warning
    assert "OCR" in warning and "可以复制文字的版本" in warning


async def test_scanned_pdf_warning_reports_the_actual_character_count():
    stamp = "P7"  # 每页只有两个字符的页码
    document = await extract(make_pdf(["image"] * 6, stamp=stamp), "扫描件.pdf")

    assert document.text != ""
    assert len(document.warnings) == 1
    warning = document.warnings[0]
    assert "扫描件" in warning and "OCR" in warning
    assert "共 6 页" in warning
    assert f"只提取到 {6 * len(stamp)} 个字符" in warning


async def test_pdf_that_is_mostly_images_is_flagged():
    # 3 页文字把页均字符数拉高了，但一多半页面其实是扫描图
    document = await extract(make_pdf(["text"] * 3 + ["image"] * 5), "混合.pdf")

    assert len(document.warnings) == 1
    warning = document.warnings[0]
    assert "共 8 页" in warning
    assert "其中 5 页几乎没有可提取的文字" in warning
    assert "OCR" in warning
    assert "Page 1 line 1" in document.text


async def test_pdf_with_some_image_pages_lists_those_pages():
    pages = ["text"] * 3 + ["image"] * 3 + ["text"] * 4
    document = await extract(make_pdf(pages), "带扫描附件.pdf")

    assert len(document.warnings) == 1
    warning = document.warnings[0]
    assert "共 10 页" in warning
    assert "其中 3 页（第 4-6 页）" in warning
    assert "很可能是扫描件" not in warning


async def test_text_pdf_with_a_blank_page_has_no_warning():
    document = await extract(make_pdf(["text"] * 5 + ["blank"] + ["text"] * 4), "招标文件.pdf")

    assert document.warnings == []


async def test_blank_docx_gets_an_empty_warning():
    document = await extract(make_docx(paragraphs=0, table=False, heading=False), "空白.docx")

    assert document.text == ""
    assert document.char_count == 0
    assert len(document.warnings) == 1
    assert "没有" in document.warnings[0] and "0 个字符" in document.warnings[0]


async def test_docx_made_of_pictures_is_flagged():
    data = make_docx(paragraphs=0, table=False, heading=False, images=3)
    document = await extract(data, "图片版.docx")

    # 图片占位标记不算内容
    assert document.text == ""
    assert len(document.warnings) == 1
    warning = document.warnings[0]
    assert "0 个字符" in warning
    assert "3 张图片" in warning
    assert "OCR" in warning


async def test_very_short_document_gets_a_warning():
    source = docx.Document()
    source.add_paragraph("招标公告")
    source.add_paragraph("详见附件。")
    buffer = io.BytesIO()
    source.save(buffer)

    document = await extract(buffer.getvalue(), "公告.docx")

    assert document.text == "招标公告\n详见附件。"
    assert document.char_count == len(document.text)
    assert len(document.warnings) == 1
    assert "只提取到 9 个字符" in document.warnings[0]


# ---------------------------------------------------------------------------
# 使用者可读的错误
# ---------------------------------------------------------------------------


async def test_pdf_content_with_docx_extension_is_rejected():
    error = await extract_error(make_pdf(["text"]), "招标文件.docx", content_type=DOCX_TYPE)

    assert "扩展名是 .docx" in error.message
    assert "PDF" in error.message and ".pdf" in error.message


async def test_docx_content_with_pdf_extension_is_rejected():
    error = await extract_error(make_docx(), "招标文件.pdf", content_type=PDF_TYPE)

    assert "扩展名是 .pdf" in error.message
    assert "Word" in error.message and ".docx" in error.message


@pytest.mark.parametrize(
    "filename, content_type, keyword",
    [("招标文件.docx", DOCX_TYPE, "Word 文档"), ("招标文件.pdf", PDF_TYPE, "PDF")],
)
async def test_plain_text_renamed_to_a_supported_extension(filename, content_type, keyword, temp_dir):
    data = "这只是一段被改了扩展名的纯文本。\n".encode("utf-8") * 50

    error = await extract_error(data, filename, content_type=content_type)

    assert f"不是有效的 {keyword}" in error.message
    assert "损坏" in error.message
    assert os.listdir(temp_dir) == []


async def test_legacy_doc_extension_tells_user_to_save_as_docx():
    data = OLE2_MAGIC + b"\x00" * 2048 + "WordDocument".encode("utf-16-le")

    error = await extract_error(data, "招标文件.doc", content_type="application/msword")

    assert ".doc" in error.message
    assert "另存为 .docx" in error.message


async def test_legacy_doc_renamed_to_docx_gets_a_readable_message():
    # 实测到的问题：老式 .doc 顶着 docx 的 Content-Type 传上来，报的是 “File is not a zip file”
    data = OLE2_MAGIC + b"\x00" * 2048 + "WordDocument".encode("utf-16-le") + b"\x00" * 4096

    error = await extract_error(data, "招标文件.docx", content_type=DOCX_TYPE)

    assert "老式 Word 格式（.doc）" in error.message
    assert "另存为 .docx" in error.message


async def test_password_protected_office_document():
    marker = "EncryptedPackage".encode("utf-16-le")
    # 标记落在很靠后的位置，而且正好跨在两个读取块的边界上
    padding = file_service._CHUNK_SIZE * 2 - len(OLE2_MAGIC) - len(marker) // 2
    data = OLE2_MAGIC + b"\x00" * padding + marker + b"\x00" * 512

    error = await extract_error(data, "招标文件.docx", content_type=DOCX_TYPE)

    assert "密码" in error.message
    assert "老式" not in error.message


async def test_other_legacy_office_file_renamed_to_docx():
    # 比如把 .xls 改名成 .docx：是 OLE2 容器，但里面没有 Word 的流
    data = OLE2_MAGIC + b"\x00" * 4096

    error = await extract_error(data, "报价表.docx", content_type=DOCX_TYPE)

    assert "老式 Office 格式" in error.message
    assert ".docx" in error.message and ".pdf" in error.message


@pytest.mark.parametrize("filename", ["报价.xlsx", "方案.txt", "文件.wps", "扫描.jpg", "标书.docm"])
async def test_unsupported_extension(filename):
    error = await extract_error(b"whatever", filename)

    assert "不支持的文件类型" in error.message
    assert ".docx" in error.message and ".pdf" in error.message
    assert "另存为 .docx" in error.message


async def test_spreadsheet_renamed_to_docx():
    data = make_zip({"[Content_Types].xml": "<Types/>", "xl/workbook.xml": "<workbook/>"})

    error = await extract_error(data, "报价表.docx", content_type=DOCX_TYPE)

    assert "不是 Word 文档" in error.message
    assert "Excel" in error.message


@pytest.mark.parametrize(
    "data",
    [
        b"PK\x03\x04" + b"garbage" * 200,
        make_docx()[:1500],
    ],
    ids=["zip 头后面全是垃圾", "下载到一半的 docx"],
)
async def test_corrupted_docx(data, temp_dir):
    error = await extract_error(data, "招标文件.docx", content_type=DOCX_TYPE)

    assert "损坏" in error.message
    assert "另存为" in error.message
    assert os.listdir(temp_dir) == []


async def test_docx_with_broken_document_xml(temp_dir):
    members = {}
    with zipfile.ZipFile(io.BytesIO(make_docx())) as package:
        for name in package.namelist():
            members[name] = package.read(name)
    members["word/document.xml"] = b"<w:document this is not xml"

    error = await extract_error(make_zip(members), "招标文件.docx")

    assert "损坏" in error.message
    assert os.listdir(temp_dir) == []


@pytest.mark.parametrize(
    "data",
    [
        b"%PDF-1.4\n" + b"this is not really a pdf " * 100,
        make_pdf(["text"] * 3)[:700],
    ],
    ids=["PDF 头后面全是垃圾", "下载到一半的 PDF"],
)
async def test_corrupted_pdf(data, temp_dir):
    error = await extract_error(data, "招标文件.pdf", content_type=PDF_TYPE)

    assert "损坏" in error.message
    assert os.listdir(temp_dir) == []


async def test_password_protected_pdf(temp_dir):
    error = await extract_error(make_encrypted_pdf("open-secret"), "招标文件.pdf")

    assert "密码" in error.message
    assert "不带密码" in error.message
    assert os.listdir(temp_dir) == []


async def test_pdf_with_only_permission_restrictions_is_still_readable():
    # 只限制了打印/复制权限、不需要密码就能打开的 PDF 照常读取
    document = await extract(make_encrypted_pdf(""), "招标文件.pdf")

    assert f"line 1: {PDF_LINE}" in document.text


async def test_empty_file():
    error = await extract_error(b"", "招标文件.docx")

    assert "空的" in error.message


async def test_docx_that_unpacks_too_large(monkeypatch):
    monkeypatch.setattr(file_service, "_MAX_DOCX_UNPACKED_BYTES", 1000)

    error = await extract_error(make_docx(), "招标文件.docx")

    assert "过于庞大" in error.message


# ---------------------------------------------------------------------------
# 大小上限
# ---------------------------------------------------------------------------


class _HugePdfStream:
    """假装是一个巨大的 PDF：数据按需产出，并记下被读走了多少字节。

    读走的字节超过 read_budget 就直接报错：实现一旦退化成“先读完再检查大小”，
    测试会立刻失败，而不是真的去读、去解析几百兆数据。
    """

    def __init__(self, total: int, read_budget: int) -> None:
        self.total = total
        self.read_budget = read_budget
        self.position = 0
        self.bytes_read = 0

    def seek(self, offset: int, whence: int = 0) -> int:
        self.position = offset
        return offset

    def read(self, size: int = -1) -> bytes:
        remaining = self.total - self.position
        size = remaining if size is None or size < 0 else min(size, remaining)
        if self.bytes_read + size > self.read_budget:
            raise AssertionError("超过大小上限之后还在继续读取上传内容")
        data = bytes(size)
        if self.position == 0 and size >= 5:
            data = b"%PDF-" + data[5:]
        self.position += size
        self.bytes_read += size
        return data


@pytest.mark.parametrize("limit", [64 * 1024, 600 * 1024], ids=["上限小于一个读取块", "上限跨多个读取块"])
async def test_oversize_upload_is_aborted_without_reading_the_rest(limit, monkeypatch, temp_dir):
    monkeypatch.setattr(settings, "max_file_size", limit)
    # 超限后立刻停手：最多只多读一个块，而不是把 200MB 读完
    stream = _HugePdfStream(
        total=200 * 1024 * 1024, read_budget=limit + file_service._CHUNK_SIZE
    )
    upload = UploadFile(stream, filename="巨大的文件.pdf")

    with pytest.raises(AppError) as caught:
        await FileService.extract_document(upload)

    assert caught.value.status_code == 400
    assert "文件太大" in caught.value.message
    assert 0 < stream.bytes_read <= limit + file_service._CHUNK_SIZE
    assert os.listdir(temp_dir) == []


async def test_declared_size_over_limit_is_rejected_before_reading(monkeypatch):
    monkeypatch.setattr(settings, "max_file_size", 1024)
    stream = _HugePdfStream(total=10 * 1024 * 1024, read_budget=0)
    upload = UploadFile(stream, filename="巨大的文件.pdf", size=10 * 1024 * 1024)

    with pytest.raises(AppError) as caught:
        await FileService.extract_document(upload)

    assert caught.value.status_code == 400
    assert "文件太大" in caught.value.message
    assert "1KB" in caught.value.message
    assert stream.bytes_read == 0


@pytest.mark.parametrize(
    "limit, shown",
    [(10 * 1024 * 1024, "10MB"), (1536 * 1024, "1.5MB"), (300 * 1024, "300KB")],
)
def test_size_limit_is_shown_in_familiar_units(limit, shown):
    assert f"超过了 {shown} 的上限" in file_service._too_large_message(limit)


@pytest.mark.parametrize(
    "pages, shown",
    [
        ([4], "第 4 页"),
        ([1, 3, 4, 5, 9], "第 1、3-5、9 页"),
        ([1, 3, 5, 7, 9, 11, 13, 15], "第 1、3、5、7、9、11 页等"),
    ],
)
def test_page_numbers_are_compacted_into_ranges(pages, shown):
    assert file_service._format_pages(pages) == shown


async def test_file_exactly_at_the_limit_is_accepted(monkeypatch):
    data = make_docx()
    monkeypatch.setattr(settings, "max_file_size", len(data))
    document = await extract(data, "招标文件.docx")
    assert f"第1段：{PARAGRAPH}" in document.text

    monkeypatch.setattr(settings, "max_file_size", len(data) - 1)
    error = await extract_error(data, "招标文件.docx")
    assert "文件太大" in error.message


# ---------------------------------------------------------------------------
# 临时文件
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "data, filename",
    [(make_docx(), "招标文件.docx"), (make_pdf(["text"] * 2), "招标文件.pdf")],
    ids=["docx", "pdf"],
)
async def test_temp_file_lives_in_system_temp_dir_and_is_removed(data, filename, temp_dir, monkeypatch):
    seen_while_parsing = []
    real_parse = file_service._parse_path

    def spying_parse(path, kind, want_images):
        seen_while_parsing.append((path, sorted(os.listdir(temp_dir))))
        return real_parse(path, kind, want_images)

    monkeypatch.setattr(file_service, "_parse_path", spying_parse)

    document = await extract(data, filename)

    assert document.text
    # 解析时文件确实落在系统临时目录里，文件名与使用者提供的名字无关
    (path, listing), = seen_while_parsing
    assert os.path.dirname(path) == str(temp_dir)
    assert listing == [os.path.basename(path)]
    assert "招标文件" not in os.path.basename(path)
    # 结束后清理干净，也不会在当前目录下建 uploads 目录
    assert os.listdir(temp_dir) == []
    assert not os.path.isdir("uploads") or os.listdir("uploads") == []


async def test_temp_file_removed_when_parser_crashes(temp_dir, monkeypatch):
    def exploding_parse(path, kind, want_images):
        assert os.path.exists(path)
        raise RuntimeError("解析器内部错误")

    monkeypatch.setattr(file_service, "_parse_path", exploding_parse)

    with pytest.raises(RuntimeError):
        await extract(make_docx(), "招标文件.docx")

    assert os.listdir(temp_dir) == []


async def test_temp_file_removed_when_request_is_cancelled(temp_dir, monkeypatch):
    entered = threading.Event()
    release = threading.Event()
    real_parse = file_service._parse_path

    def blocked_parse(path, kind, want_images):
        entered.set()
        release.wait(5)
        return real_parse(path, kind, want_images)

    monkeypatch.setattr(file_service, "_parse_path", blocked_parse)
    loop = asyncio.get_running_loop()

    task = asyncio.create_task(extract(make_docx(), "招标文件.docx"))
    try:
        assert await asyncio.to_thread(entered.wait, 5)
        assert len(os.listdir(temp_dir)) == 1  # 解析进行中，临时文件还在

        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
    finally:
        release.set()

    # 解析线程是串行的：排在后面的空任务跑完，说明被取消的那次解析已经收尾
    await loop.run_in_executor(file_service._PARSE_EXECUTOR, lambda: None)
    assert os.listdir(temp_dir) == []


# ---------------------------------------------------------------------------
# 事件循环
# ---------------------------------------------------------------------------


async def test_event_loop_keeps_running_while_parsing(monkeypatch):
    block_seconds = 1.0
    parse_threads = []
    real_parse = file_service._parse_path

    def slow_parse(path, kind, want_images):
        parse_threads.append(threading.get_ident())
        # 模拟一个要解析很久的大文件：解析要是跑在事件循环线程上，心跳就会被卡住这么久
        time.sleep(block_seconds)
        return real_parse(path, kind, want_images)

    monkeypatch.setattr(file_service, "_parse_path", slow_parse)

    gaps = []
    stop = asyncio.Event()

    async def heartbeat():
        last = time.perf_counter()
        while not stop.is_set():
            await asyncio.sleep(0.01)
            now = time.perf_counter()
            gaps.append(now - last)
            last = now

    big = make_docx(paragraphs=600, table_rows=[("评分项", "分值", "评分标准")] * 60)
    beat = asyncio.create_task(heartbeat())
    await asyncio.sleep(0.05)
    gaps.clear()

    started = time.perf_counter()
    document = await extract(big, "大文档.docx")
    elapsed = time.perf_counter() - started
    stop.set()
    await beat

    assert document.char_count > 20000
    assert elapsed >= block_seconds
    assert parse_threads and parse_threads[0] != threading.get_ident()
    # 解析期间心跳一直在跳，没有哪一次被卡住接近解析耗时那么久
    assert len(gaps) >= 10
    assert max(gaps) < block_seconds / 2, f"事件循环被阻塞了 {max(gaps):.3f}s"


async def test_concurrent_uploads_all_succeed(temp_dir):
    uploads = [
        extract(make_docx(paragraphs=20), "甲.docx"),
        extract(make_pdf(["text"] * 3), "乙.pdf"),
        extract(make_docx(paragraphs=5), "丙.docx"),
        extract(make_pdf(["image"] * 2), "丁.pdf"),
    ]

    first, second, third, fourth = await asyncio.gather(*uploads)

    assert f"第20段：{PARAGRAPH}" in first.text and first.filename == "甲.docx"
    assert "--- 第 3 页 ---" in second.text and second.warnings == []
    assert f"第5段：{PARAGRAPH}" in third.text and f"第6段：{PARAGRAPH}" not in third.text
    assert fourth.text == "" and "扫描件" in fourth.warnings[0]
    assert os.listdir(temp_dir) == []


# ---------------------------------------------------------------------------
# 备用解析器
# ---------------------------------------------------------------------------


async def test_docx_falls_back_to_python_docx(monkeypatch):
    def broken_docx2python(*args, **kwargs):
        raise ValueError("docx2python 不认识这个文档")

    monkeypatch.setattr(file_service, "docx2python", broken_docx2python)

    document = await extract(make_docx(), "招标文件.docx")

    lines = document.text.splitlines()
    assert f"第8段：{PARAGRAPH}" in lines
    assert "技术方案 | 30 | 方案完整、针对性强" in lines
    # 备用解析器同样按文档顺序输出，表格不会被挪到文末
    assert document.text.index("[表格结束]") < document.text.index("表格后的段落。")
    assert document.warnings == []


async def test_pdf_falls_back_to_pymupdf_and_says_so(monkeypatch, temp_dir):
    def broken_pdfplumber(path):
        raise ValueError("pdfplumber 读不了这个文件")

    monkeypatch.setattr(file_service, "_parse_pdf_pdfplumber", broken_pdfplumber)

    document = await extract(make_pdf(["text"] * 2), "招标文件.pdf")

    assert f"Page 2 line 20: {PDF_LINE}" in document.text
    assert "--- 第 2 页 ---" in document.text
    assert len(document.warnings) == 1
    assert "内容可能不完整" in document.warnings[0]
    assert os.listdir(temp_dir) == []


@pytest.mark.filterwarnings("ignore:Advanced encoding")  # PyPDF2 不认识测试 PDF 里的中文字体编码
async def test_basic_parsers_when_advanced_libraries_are_missing(monkeypatch, temp_dir):
    monkeypatch.setattr(file_service, "HAS_ADVANCED_LIBS", False)

    word = await extract(make_docx(), "招标文件.docx")
    assert f"第1段：{PARAGRAPH}" in word.text.splitlines()
    assert "商务报价 | 40 | 按评标基准价计算" in word.text.splitlines()

    pdf = await extract(make_pdf(["text"] * 2, lines=5), "招标文件.pdf")
    assert "--- 第 2 页 ---" in pdf.text
    assert "Page 2 line 5" in pdf.text

    scanned = await extract(make_pdf(["image"] * 3), "扫描件.pdf")
    assert scanned.text == "" and "扫描件" in scanned.warnings[0]

    error = await extract_error(make_encrypted_pdf("open-secret"), "加密.pdf")
    assert "密码" in error.message
    assert os.listdir(temp_dir) == []


# ---------------------------------------------------------------------------
# 旧入口
# ---------------------------------------------------------------------------


async def test_process_uploaded_file_still_returns_plain_text():
    data = make_docx()

    text = await FileService.process_uploaded_file(make_upload(data, "招标文件.docx", DOCX_TYPE))

    assert isinstance(text, str)
    assert text == (await extract(data, "招标文件.docx")).text


async def test_process_uploaded_file_refuses_documents_without_any_text():
    upload = make_upload(make_pdf(["image"] * 4), "扫描件.pdf", PDF_TYPE)

    with pytest.raises(AppError) as caught:
        await FileService.process_uploaded_file(upload)

    assert caught.value.status_code == 400
    assert "扫描件" in caught.value.message


def test_is_supported_document_keeps_its_contract():
    assert FileService.is_supported_document(DOCX_TYPE)
    assert FileService.is_supported_document(PDF_TYPE)
    assert not FileService.is_supported_document("application/msword")
    assert not FileService.is_supported_document(None)


async def test_extracting_from_a_local_path(tmp_path):
    docx_path = tmp_path / "招标文件.docx"
    docx_path.write_bytes(make_docx())
    pdf_path = tmp_path / "招标文件.pdf"
    pdf_path.write_bytes(make_pdf(["text"]))

    assert f"第1段：{PARAGRAPH}" in await FileService.extract_text_from_docx(str(docx_path))
    assert f"Page 1 line 1: {PDF_LINE}" in await FileService.extract_text_from_pdf(str(pdf_path))
    # 抽取完不占着文件
    docx_path.unlink()
    pdf_path.unlink()


async def test_upload_can_be_read_again_afterwards():
    data = make_docx()
    upload = make_upload(data, "招标文件.docx")

    await FileService.extract_document(upload)

    assert await upload.read() == data


# ---------------------------------------------------------------------------
# 可选的图片外传
# ---------------------------------------------------------------------------


async def test_no_external_request_unless_image_upload_is_configured(monkeypatch):
    import aiohttp

    sessions = []

    def forbidden_session(*args, **kwargs):
        sessions.append(args)
        raise AssertionError("未配置 IMAGE_UPLOAD_URL 时不应该创建任何网络会话")

    monkeypatch.setattr(aiohttp, "ClientSession", forbidden_session)

    document = await extract(make_docx(images=2), "带图.docx")

    assert sessions == []
    # 图片位置只留下占位标记，没有任何外链
    assert "----media/image1.png----" in document.text
    assert "图片引用" not in document.text
    assert await FileService.upload_image_to_server(b"image-bytes", "a.jpg") is None
    assert sessions == []


async def test_images_are_linked_when_image_upload_is_configured(monkeypatch):
    calls = []

    async def fake_upload(image_data: bytes, filename: str):
        calls.append((image_data, filename))
        return f"http://image-host.invalid/files/{filename}"

    monkeypatch.setattr(FileService, "IMAGE_UPLOAD_URL", "http://image-host.invalid/upload")
    monkeypatch.setattr(FileService, "upload_image_to_server", staticmethod(fake_upload))

    document = await extract(make_docx(images=2), "带图.docx")

    # 每张图片传一次，传上去的就是文档里对应位置的那张图
    assert calls == [(_png_bytes(0), "docx_img1.png"), (_png_bytes(1), "docx_img2.png")]
    assert "----media/" not in document.text
    assert document.text.index("[图片1]") < document.text.index("[图片2]")
    assert document.text.endswith(
        "--- 图片引用 ---\n"
        "[图片1]: http://image-host.invalid/files/docx_img1.png\n"
        "[图片2]: http://image-host.invalid/files/docx_img2.png"
    )
    assert document.char_count == len(document.text)


async def test_image_used_twice_is_uploaded_once(monkeypatch):
    calls = []

    async def fake_upload(image_data: bytes, filename: str):
        calls.append(filename)
        return f"http://image-host.invalid/files/{filename}"

    monkeypatch.setattr(FileService, "IMAGE_UPLOAD_URL", "http://image-host.invalid/upload")
    monkeypatch.setattr(FileService, "upload_image_to_server", staticmethod(fake_upload))

    source = docx.Document()
    for index in range(6):
        source.add_paragraph(f"第{index + 1}段：{PARAGRAPH}")
    source.add_picture(io.BytesIO(_png_bytes()))
    source.add_paragraph("两处用的是同一张图。")
    source.add_picture(io.BytesIO(_png_bytes()))
    buffer = io.BytesIO()
    source.save(buffer)

    document = await extract(buffer.getvalue(), "带图.docx")

    assert calls == ["docx_img1.png"]
    assert document.text.count("[图片1]") == 3  # 正文两处，加文末的引用
    assert "----media/" not in document.text


async def test_failed_image_upload_keeps_the_placeholder(monkeypatch):
    async def failing_upload(image_data: bytes, filename: str):
        return None

    monkeypatch.setattr(FileService, "IMAGE_UPLOAD_URL", "http://image-host.invalid/upload")
    monkeypatch.setattr(FileService, "upload_image_to_server", staticmethod(failing_upload))

    document = await extract(make_docx(images=1), "带图.docx")

    assert "----media/image1.png----" in document.text
    assert "图片引用" not in document.text


# ---------------------------------------------------------------------------
# /api/document/upload
# ---------------------------------------------------------------------------


@pytest.fixture()
def client():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.routers import document

    app = FastAPI()
    app.include_router(document.router)
    with TestClient(app) as test_client:
        yield test_client


def _post(client, data: bytes, filename: str, content_type: str):
    return client.post("/api/document/upload", files={"file": (filename, data, content_type)})


def test_upload_endpoint_success_shape(client):
    response = _post(client, make_docx(), "招标文件.docx", DOCX_TYPE)

    assert response.status_code == 200
    body = response.json()
    assert body["success"] is True
    assert body["message"] == "文件 招标文件.docx 上传成功"
    assert f"第1段：{PARAGRAPH}" in body["file_content"]
    assert {"success", "message", "file_content"} <= set(body)


def test_upload_endpoint_handles_uploads_spooled_to_disk(client):
    # 超过 1MB 的上传会被 Starlette 从内存转存到磁盘上的临时文件，读取走的是另一条路径
    noise = random.Random(0).randbytes(700 * 600 * 3)  # 压不动的噪点图，把文档撑到 1MB 以上
    picture = fitz.Pixmap(fitz.csRGB, 700, 600, noise, False).tobytes("png")
    source = docx.Document()
    for index in range(8):
        source.add_paragraph(f"第{index + 1}段：{PARAGRAPH}")
    source.add_picture(io.BytesIO(picture))
    buffer = io.BytesIO()
    source.save(buffer)
    data = buffer.getvalue()
    assert len(data) > 1024 * 1024

    response = _post(client, data, "带大图的招标文件.docx", DOCX_TYPE)

    body = response.json()
    assert response.status_code == 200
    assert body["success"] is True
    assert f"第8段：{PARAGRAPH}" in body["file_content"]


def test_upload_endpoint_does_not_depend_on_content_type(client):
    response = _post(client, make_pdf(["text"] * 2), "招标文件.pdf", "application/octet-stream")

    body = response.json()
    assert response.status_code == 200
    assert body["success"] is True
    assert "--- 第 2 页 ---" in body["file_content"]


def test_upload_endpoint_reports_failure_with_http_200(client):
    data = OLE2_MAGIC + b"\x00" * 4096 + "WordDocument".encode("utf-16-le")

    response = _post(client, data, "招标文件.docx", DOCX_TYPE)

    assert response.status_code == 200
    body = response.json()
    assert body["success"] is False
    assert body["file_content"] is None
    assert "另存为 .docx" in body["message"]
    assert "zip" not in body["message"]


def test_upload_endpoint_rejects_oversize_files(client, monkeypatch):
    data = make_docx()
    monkeypatch.setattr(settings, "max_file_size", len(data) - 1)

    response = _post(client, data, "招标文件.docx", DOCX_TYPE)

    assert response.status_code == 200
    body = response.json()
    assert body["success"] is False
    assert "文件太大" in body["message"]


def test_upload_endpoint_appends_warnings_to_the_message(client):
    source = docx.Document()
    source.add_paragraph("招标公告：详见附件。")
    buffer = io.BytesIO()
    source.save(buffer)

    response = _post(client, buffer.getvalue(), "公告.docx", DOCX_TYPE)

    body = response.json()
    assert response.status_code == 200
    assert body["success"] is True
    assert body["file_content"] == "招标公告：详见附件。"
    assert body["message"].startswith("文件 公告.docx 上传成功，但请注意：")
    assert "只提取到 10 个字符" in body["message"]


def test_upload_endpoint_treats_scanned_pdf_as_failure(client):
    response = _post(client, make_pdf(["image"] * 5), "盖章扫描件.pdf", PDF_TYPE)

    assert response.status_code == 200
    body = response.json()
    assert body["success"] is False
    assert body["file_content"] is None
    assert "扫描件" in body["message"] and "OCR" in body["message"]
    assert "上传成功" not in body["message"]


def test_upload_endpoint_hides_unexpected_errors(client, monkeypatch):
    async def exploding_extract(file):
        raise RuntimeError("internal detail: boom")

    monkeypatch.setattr(FileService, "extract_document", staticmethod(exploding_extract))

    response = _post(client, make_docx(), "招标文件.docx", DOCX_TYPE)

    assert response.status_code == 200
    body = response.json()
    assert body["success"] is False
    assert "boom" not in body["message"]
    assert "文件处理失败" in body["message"]
