"""Word 导出的测试：Markdown 渲染不丢字、真表格、规范排版确定化、旧多份路径仍可用。"""
from __future__ import annotations

import base64
import io
import random
import zipfile
from datetime import datetime, timezone

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from PIL import Image

from app.models.schemas import WordExportOutlineItem, WordExportRequest
from app.services import word_export_service as wes
from app.services.word_export_service import WordExportService

_TWIP = 635  # 1 缇对应的 EMU，页边距写入时按缇取整


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------


def _export_bytes(content: str = "正文", *, outline=None, **fields) -> bytes:
    if outline is None:
        outline = [WordExportOutlineItem(id="1", title="技术方案", content=content)]
    request = WordExportRequest(outline=outline, **fields)
    buffer, headers = WordExportService.export_outline(request)
    assert isinstance(buffer, io.BytesIO)
    assert headers["Content-Disposition"].startswith("attachment; filename*=UTF-8''")
    return buffer.getvalue()


def _export(content: str = "正文", **kwargs):
    """导出后用 python-docx 读回。"""
    return docx.Document(io.BytesIO(_export_bytes(content, **kwargs)))


def _all_text(document) -> str:
    """正文区的全部文字（含表格单元格），按文档顺序拼接。"""
    return "".join(node.text or "" for node in document.element.body.iter(qn("w:t")))


def _para(document, text: str):
    """正文里文字恰为 text 的段落（不含表格内的段落）。"""
    matches = [p for p in document.paragraphs if p.text == text]
    assert matches, f"没有找到段落：{text!r}；现有段落：{[p.text for p in document.paragraphs]}"
    return matches[0]


def _texts(document) -> list[str]:
    return [p.text for p in document.paragraphs]


def _east_asia_font(run_or_style) -> str | None:
    rpr = run_or_style.element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get(qn("w:eastAsia"))


def _cjk(text: str) -> str:
    return "".join(ch for ch in text if "\u4e00" <= ch <= "\u9fff")


def _span_text(spans) -> str:
    return "".join("\n" if kind == "br" else text for kind, text, _bold, _italic in spans)


# ---------------------------------------------------------------------------
# 缺陷 1：以 * 或 - 开头但不是列表的段落不能丢
# ---------------------------------------------------------------------------


def test_paragraphs_starting_with_bold_or_dash_are_kept():
    document = _export(
        "**（1）系统架构**：采用分层设计，各层之间松耦合。\n"
        "-10℃～50℃ 环境下可稳定运行。\n"
        "*注：以上参数为实测值。\n"
        "+86 为国际区号。\n"
    )

    bold_lead = _para(document, "（1）系统架构：采用分层设计，各层之间松耦合。")
    assert bold_lead.runs[0].text == "（1）系统架构" and bold_lead.runs[0].bold
    assert not bold_lead.runs[1].bold

    _para(document, "-10℃～50℃ 环境下可稳定运行。")
    _para(document, "*注：以上参数为实测值。")
    _para(document, "+86 为国际区号。")


# ---------------------------------------------------------------------------
# 缺陷 2：表格
# ---------------------------------------------------------------------------

_TABLE_MD = (
    "设备清单如下：\n"
    "\n"
    "| 名称 | 数量 | 备注 |\n"
    "|:-----|:----:|-----:|\n"
    "| 服务器 | 2 | |\n"
    "| 交换机 |  | 核心<br>冗余 |\n"
    "| **防火墙** | 1 | 甲 \\| 乙 | 多出一格 |\n"
    "| 仅一格 |\n"
    "\n"
    "表后正文。\n"
)


def test_markdown_table_becomes_real_word_table():
    document = _export(_TABLE_MD)

    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.style.name == "Table Grid"  # 带边框
    assert len(table.rows) == 5 and len(table.columns) == 3
    assert [[cell.text for cell in row.cells] for row in table.rows] == [
        ["名称", "数量", "备注"],
        ["服务器", "2", ""],                    # 行尾空单元格
        ["交换机", "", "核心\n冗余"],            # 中间空单元格不错位；<br> 变成单元格内换行
        ["防火墙", "1", "甲 | 乙 | 多出一格"],   # 多出的单元格并入最后一列，不丢字
        ["仅一格", "", ""],                      # 不足的按表头补齐
    ]

    for cell in table.rows[0].cells:
        runs = cell.paragraphs[0].runs
        assert runs and all(run.bold for run in runs), "表头应加粗"
    body_run = table.rows[1].cells[0].paragraphs[0].runs[0]
    assert not body_run.bold
    assert _east_asia_font(body_run) == document.styles["Normal"].font.name == "宋体"
    assert table.rows[3].cells[0].paragraphs[0].runs[0].bold, "单元格内的 **加粗** 应生效"
    # 分隔行里的对齐标记生效
    assert table.rows[1].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert table.rows[1].cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    # 表格内容不再以「a | b | c」的文本段落出现
    assert _texts(document).count("设备清单如下：") == 1
    assert not any("服务器" in text or "|" in text for text in _texts(document))
    # 表格上下与正文之间留有间距
    assert _para(document, "设备清单如下：").paragraph_format.space_after == Pt(6)
    assert _para(document, "表后正文。").paragraph_format.space_before == Pt(6)


def test_table_layout_details():
    document = _export(_TABLE_MD)
    table = document.tables[0]
    section = document.sections[0]

    # 列宽按内容分配并占满版心：备注列内容最长，数量列最短
    widths = [column.width for column in table.columns]
    assert widths[2] > widths[0] > widths[1] > 0
    text_width = section.page_width - section.left_margin - section.right_margin
    assert abs(sum(widths) - text_width) <= 2 * _TWIP
    for row in table.rows:
        assert [cell.width for cell in row.cells] == widths
        # 行不跨页拆开
        assert row._tr.trPr.find(qn("w:cantSplit")) is not None
    # 只有表头行在跨页时重复
    assert [row._tr.trPr.find(qn("w:tblHeader")) is not None for row in table.rows] == [
        True, False, False, False, False,
    ]


def test_column_widths_always_fill_text_width():
    total, unit = 8311, 120
    # 内容都放得下：等比放大，短内容的两列等宽
    assert wes._table_column_widths([[4, 4]], total, unit) == [4155, 4156]
    # 放不下：短列（序号）拿到所需宽度，两个长列平分其余
    narrow, wide_a, wide_b = wes._table_column_widths([[4, 100, 100], [2, 80, 90]], total, unit)
    assert narrow == 4 * unit + 336 and abs(wide_a - wide_b) <= 1
    for cells in ([[1] * 63], [[0, 0, 0]], [[500, 1, 1, 1]], [[7, 13, 29, 31, 37]]):
        widths = wes._table_column_widths(cells, total, unit)
        assert sum(widths) == total and min(widths) > 0


def test_adjacent_tables_stay_separate():
    document = _export("| 甲 |\n|---|\n| 1 |\n\n| 乙 |\n|---|\n| 2 |\n")

    assert [[c.text for c in row.cells] for t in document.tables for row in t.rows] == [
        ["甲"], ["1"], ["乙"], ["2"],
    ]
    # 两张表之间必须隔着段落，否则 Word 打开时会把它们并成一张
    tags = [child.tag for child in document.element.body]
    first = tags.index(qn("w:tbl"))
    assert tags[first + 1] == qn("w:p") and tags[first + 2] == qn("w:tbl")


def test_document_never_ends_with_a_table():
    document = _export("| 甲 | 乙 |\n|---|---|\n| 1 | 2 |")

    tags = [child.tag for child in document.element.body]
    assert tags[-3:] == [qn("w:tbl"), qn("w:p"), qn("w:sectPr")]


def test_table_wider_than_word_allows_falls_back_to_text():
    width = 64  # Word 表格最多 63 列
    source = "\n".join(
        [
            "|" + "|".join(f"列{n}" for n in range(width)) + "|",
            "|" + "|".join("---" for _ in range(width)) + "|",
            "|" + "|".join(f"值{n}" for n in range(width)) + "|",
        ]
    )
    document = _export(source)

    assert document.tables == []
    text = _all_text(document)
    assert all(f"列{n}" in text and f"值{n}" in text for n in range(width))


def test_pipe_text_without_separator_row_stays_paragraph():
    document = _export("阶段一 | 阶段二 | 阶段三\n命令 ls | grep 关键字\n\n甲 | 乙\n---\n结尾")

    assert document.tables == []
    _para(document, "阶段一 | 阶段二 | 阶段三")
    _para(document, "命令 ls | grep 关键字")
    _para(document, "甲 | 乙")
    assert "---" not in _texts(document)


# ---------------------------------------------------------------------------
# 缺陷 3：正文里的 # 标题不能占用 Word 标题层级
# ---------------------------------------------------------------------------


def test_body_headings_are_bold_paragraphs_while_chapters_use_heading_styles():
    leaf = WordExportOutlineItem(
        id="1.1.1",
        title="部署方案",
        content="# 一号小标题\n正文甲。\n## 二号小标题\n###### 六号小标题\n##漏了空格的小标题\n#1 机组不是标题\n",
    )
    deep = WordExportOutlineItem(id="1.1.2.1", title="四级章节", content="正文乙。")
    outline = [
        WordExportOutlineItem(
            id="1",
            title="总体方案",
            children=[
                WordExportOutlineItem(
                    id="1.1",
                    title="技术路线",
                    children=[leaf, WordExportOutlineItem(id="1.1.2", title="实施步骤", children=[deep])],
                )
            ],
        )
    ]
    document = _export(outline=outline, project_name="示例项目")

    assert _para(document, "1 总体方案").style.name == "Heading 1"
    assert _para(document, "1.1 技术路线").style.name == "Heading 2"
    assert _para(document, "1.1.1 部署方案").style.name == "Heading 3"
    assert _para(document, "1.1.2 实施步骤").style.name == "Heading 3"

    for text in ("一号小标题", "二号小标题", "六号小标题", "漏了空格的小标题"):
        para = _para(document, text)
        assert para.style.name == "Normal", f"{text} 不应使用标题样式"
        assert para.runs and all(run.bold for run in para.runs)
        assert not para.paragraph_format.first_line_indent

    # 文档里使用标题样式的只有章节标题
    assert [p.text for p in document.paragraphs if p.style.name.startswith("Heading")] == [
        "1 总体方案", "1.1 技术路线", "1.1.1 部署方案", "1.1.2 实施步骤",
    ]
    # 四级及更深的章节用加粗段落
    deep_title = _para(document, "1.1.2.1 四级章节")
    assert deep_title.style.name == "Normal" and deep_title.runs[0].bold
    # 「#1 机组」这类正文原样保留
    assert not _para(document, "#1 机组不是标题").runs[0].bold


def test_chapter_heading_text_is_tidied():
    outline = [
        WordExportOutlineItem(id="", title="没有编号的章节", content="正文甲。"),
        WordExportOutlineItem(id="2", title="标题里\n有换行", content="正文乙。"),
        WordExportOutlineItem(id="", title="", content="无标题节点的正文。"),
    ]
    document = _export(outline=outline)

    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["没有编号的章节", "2 标题里 有换行"]  # 不留前导空格、不留空标题
    _para(document, "无标题节点的正文。")


# ---------------------------------------------------------------------------
# 列表、代码围栏、引用、分隔线
# ---------------------------------------------------------------------------


def test_nested_lists_are_indented_and_keep_numbers():
    document = _export(
        "1. 第一步\n"
        "   - 子项甲\n"
        "     - 孙项\n"
        "   - 子项乙\n"
        "\n"
        "     子项乙的补充说明\n"
        "2. 第二步\n"
        "10) 第十步\n"
        "* 星号列表项\n"
    )

    def left(text: str) -> int:
        return _para(document, text).paragraph_format.left_indent or 0

    assert left("1. 第一步") < left("• 子项甲") < left("• 孙项")
    assert left("• 子项乙") == left("• 子项甲")
    assert left("子项乙的补充说明") == left("• 子项乙")  # 续行跟随所属列表项
    assert left("2. 第二步") == left("1. 第一步")
    # 有序列表保留原编号
    assert [t for t in _texts(document) if t[:1].isdigit()][-3:] == ["1. 第一步", "2. 第二步", "10) 第十步"]
    _para(document, "• 星号列表项")


def test_code_fence_quote_and_horizontal_rule():
    document = _export(
        "说明如下：\n"
        "```python\n"
        "def 计算(x):\n"
        "    return x * 2  # **不是加粗**\n"
        "```\n"
        "> 引用内容，含**重点**。\n"
        "\n"
        "---\n"
        "***\n"
        "结尾段落\n"
    )

    code_font = wes._StandardProfile.code_font
    for text in ("def 计算(x):", "    return x * 2  # **不是加粗**"):
        para = _para(document, text)  # 逐行输出，缩进与星号原样保留
        assert para.runs[0].font.name == code_font
        assert not para.paragraph_format.first_line_indent

    quote = _para(document, "引用内容，含重点。")
    assert quote.paragraph_format.left_indent > 0
    assert not quote.paragraph_format.first_line_indent
    assert [run.text for run in quote.runs if run.bold] == ["重点"]

    texts = _texts(document)
    assert "---" not in texts and "***" not in texts
    assert not any(text.startswith("```") for text in texts)
    _para(document, "结尾段落")


def test_pathological_nesting_does_not_break_export():
    document = _export("> " * 3000 + "深层引用\n" + "\n".join(" " * (2 * n) + f"- 第{n}层" for n in range(300)))

    text = _all_text(document)
    assert "深层引用" in text and "第299层" in text
    # 缩进有上限，文字不会被挤出版心
    section = document.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    assert max(p.paragraph_format.left_indent or 0 for p in document.paragraphs) < text_width / 2


def test_fence_wrapping_whole_answer_is_unwrapped():
    wrapped = _export("```markdown\n## 小标题\n正文**加粗**。\n```\n")
    assert _para(wrapped, "小标题").runs[0].bold
    assert _para(wrapped, "正文加粗。").paragraph_format.first_line_indent == Pt(24)

    # 没有闭合的围栏标记只丢标记本身，后面的正文照常渲染
    dangling = _export("前文。\n```\n后文第一行\n后文第二行\n")
    assert _texts(dangling)[-3:] == ["前文。", "后文第一行", "后文第二行"]


# ---------------------------------------------------------------------------
# 行内格式
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    [
        "面积 2*3 平方米，另有 4*5 备用",
        "*注：带*号的为必填项，带*的为关键指标",
        "支持 *.jpg、*.png 格式",
        "联系电话 138****1234、139****5678",
        "幂运算 2**3 与 4**5",
        "乘号前后有空格：a * b * c",
        "**没有闭合的加粗",
        "没有闭合的`行内代码",
        "日期：____年__月__日",
        "路径 C:\\Program Files\\投标工具",
        "阈值 a<b 且 b>c，R&D 部门",
    ],
)
def test_stray_markers_are_kept_verbatim(text):
    spans = wes._parse_inline(text)
    assert _span_text(spans) == text
    assert not any(bold or italic for _kind, _text, bold, italic in spans)


@pytest.mark.parametrize(
    "text, expected",
    [
        ("**注意：**正文", [("text", "注意：", True, False), ("text", "正文", False, False)]),
        ("这是**加粗**文字", [("text", "这是", False, False), ("text", "加粗", True, False), ("text", "文字", False, False)]),
        ("这是 *斜体* 文字", [("text", "这是 ", False, False), ("text", "斜体", False, True), ("text", " 文字", False, False)]),
        ("***又粗又斜***", [("text", "又粗又斜", True, True)]),
        ("**带*号的为必填项**", [("text", "带*号的为必填项", True, False)]),
        ("用`a*b*c`表示", [("text", "用", False, False), ("code", "a*b*c", False, False), ("text", "表示", False, False)]),
        ("7\\*24 小时", [("text", "7*24 小时", False, False)]),
        ("上一行<br/>下一行", [("text", "上一行", False, False), ("br", "", False, False), ("text", "下一行", False, False)]),
        ("甲&nbsp;乙&lt;丙&unknown;", [("text", "甲\u00a0乙<丙&unknown;", False, False)]),
    ],
)
def test_inline_formatting(text, expected):
    assert wes._parse_inline(text) == expected


def test_inline_formatting_reaches_document():
    document = _export("其中**加粗**、 *斜体* 与`代码`并存，2*3 不变。")
    para = _para(document, "其中加粗、 斜体 与代码并存，2*3 不变。")
    assert [run.text for run in para.runs if run.bold] == ["加粗"]
    assert [run.text for run in para.runs if run.italic] == ["斜体"]
    assert [run.text for run in para.runs if run.font.name == wes._StandardProfile.code_font] == ["代码"]


# ---------------------------------------------------------------------------
# 缺陷 4：XML 非法字符
# ---------------------------------------------------------------------------


def test_clean_text_removes_illegal_xml_characters():
    dirty = "甲\x00乙\x08丙\x1f丁\ufffe戊\uffff己\ud800庚\t辛\r\n壬\r癸"
    assert wes._clean_text(dirty) == "甲乙丙丁戊己庚\t辛\n壬\n癸"
    assert wes._clean_text(None) == ""


def test_control_characters_do_not_break_export():
    dirty = "\x00\x01\x08\x0e\x1f\x7f\ufffe\uffff"
    outline = [
        WordExportOutlineItem(
            id="1",
            title=f"安全{dirty}方案",
            content=(
                f"首段{dirty}文字保留。\n"
                f"- 列表{dirty}项\n"
                f"\n"
                f"| 表{dirty}头 | 列 |\n|---|---|\n| 单元{dirty}格 | 值 |\n"
                f"```\n代码{dirty}行\n```\n"
            ),
            image_title=f"图{dirty}注",
        )
    ]
    document = _export(
        outline=outline,
        project_name=f"某{dirty}项目",
        project_overview=f"概述{dirty}内容",
    )

    text = _all_text(document)
    for expected in ("某项目", "概述内容", "1 安全方案", "首段文字保留。", "列表项", "表头", "单元格", "代码行"):
        assert expected in text
    assert not any(ch in text for ch in dirty)


# ---------------------------------------------------------------------------
# 缺陷 5：默认排版确定化
# ---------------------------------------------------------------------------


def _layout_fingerprint(document) -> dict:
    styles = document.styles
    section = document.sections[0]
    fingerprint = {
        "margins": (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin),
        "page": (section.page_width, section.page_height),
    }
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        style = styles[name]
        fingerprint[name] = (
            style.font.name,
            _east_asia_font(style),
            style.font.size,
            style.font.bold,
            style.paragraph_format.line_spacing,
        )
    return fingerprint


def test_default_export_uses_fixed_standard_layout():
    content = "第一段正文。\n\n第二段正文。"
    first_bytes = _export_bytes(content, project_name="示例项目")
    second_bytes = _export_bytes(content, project_name="示例项目")
    first = docx.Document(io.BytesIO(first_bytes))
    second = docx.Document(io.BytesIO(second_bytes))

    # 两次导出的版式完全一致
    assert _layout_fingerprint(first) == _layout_fingerprint(second)
    for part in ("word/document.xml", "word/styles.xml", "word/settings.xml", "word/footer1.xml"):
        assert zipfile.ZipFile(io.BytesIO(first_bytes)).read(part) == zipfile.ZipFile(
            io.BytesIO(second_bytes)
        ).read(part), f"{part} 两次导出不一致"

    styles = first.styles
    # 正文：宋体小四、1.5 倍行距
    normal = styles["Normal"]
    assert normal.font.name == "宋体" and _east_asia_font(normal) == "宋体"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5
    # 一二三级标题：黑体，三号 / 四号 / 小四，加粗，黑色
    for name, size in (("Heading 1", Pt(16)), ("Heading 2", Pt(14)), ("Heading 3", Pt(12))):
        style = styles[name]
        assert style.font.name == "黑体" and _east_asia_font(style) == "黑体"
        assert style.font.size == size and style.font.bold
        assert str(style.font.color.rgb) == "000000"
        # 主题字体属性必须去掉，否则 Word 会忽略上面指定的字体
        assert not [key for key in style.element.rPr.rFonts.attrib if key.endswith("Theme")]

    # 页边距上下 2.54cm、左右 3.17cm，A4 纸
    section = first.sections[0]
    assert abs(section.top_margin - Cm(2.54)) <= _TWIP and abs(section.bottom_margin - Cm(2.54)) <= _TWIP
    assert abs(section.left_margin - Cm(3.17)) <= _TWIP and abs(section.right_margin - Cm(3.17)) <= _TWIP
    assert abs(section.page_width - Mm(210)) <= _TWIP and abs(section.page_height - Mm(297)) <= _TWIP

    # 正文段落首行缩进 2 字符
    body = _para(first, "第一段正文。")
    assert body.style.name == "Normal"
    assert body.paragraph_format.first_line_indent == Pt(24)
    assert body._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
    assert _east_asia_font(body.runs[0]) == "宋体"
    heading_run = _para(first, "1 技术方案").runs[0]
    assert heading_run.font.name == "黑体" and _east_asia_font(heading_run) == "黑体"


def test_default_export_writes_no_fake_author():
    before = datetime.now(timezone.utc).replace(tzinfo=None)
    props_list = [_export("正文", project_name="示例项目").core_properties for _ in range(3)]
    after = datetime.now(timezone.utc).replace(tzinfo=None)

    # 旧多份路径的候选姓名 / 部门（该路径若日后移除，这里退化为只查模板占位信息）
    fake_words = {"python-docx", *getattr(wes, "_FAKE_AUTHORS", ()), *getattr(wes, "_FAKE_COMPANIES", ())}
    for props in props_list:
        assert props.author == "" and props.last_modified_by == ""
        assert props.comments == "" and props.category == "" and props.keywords == ""
        assert props.title == "示例项目"
        assert not any(
            word in value
            for word in fake_words
            for value in (props.author, props.last_modified_by, props.comments, props.category, props.title)
        )
        # 创建、修改时间是当前时间（文档属性按 UTC 存储，精确到秒）
        for stamp in (props.created, props.modified):
            stamp = stamp.replace(tzinfo=None)
            assert (before - stamp).total_seconds() <= 1 and (stamp - after).total_seconds() <= 1


def test_footer_has_centered_page_number_field():
    footer = _export("正文").sections[0].footer
    para = footer.paragraphs[0]
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    instructions = [node.text.strip() for node in footer._element.iter(qn("w:instrText"))]
    assert instructions == ["PAGE"]
    marks = [node.get(qn("w:fldCharType")) for node in footer._element.iter(qn("w:fldChar"))]
    assert marks == ["begin", "separate", "end"]


def test_overview_section_only_when_present():
    for blank in (None, "", "  \n "):
        document = _export("正文", project_name="示例项目", project_overview=blank)
        assert "项目概述" not in _texts(document)
        assert _texts(document)[0] == "示例项目"

    document = _export("正文", project_overview="本项目位于**某市**。\n建设周期 12 个月。")
    assert _para(document, "项目概述").style.name == "Heading 1"
    assert _para(document, "本项目位于某市。").paragraph_format.first_line_indent == Pt(24)
    _para(document, "建设周期 12 个月。")
    # 没给项目名称时用默认标题
    assert _texts(document)[0] == "投标技术文件"


def test_include_toc_inserts_toc_field_before_body():
    def toc_instructions(document) -> list[str]:
        return [
            node.text.strip()
            for node in document.element.body.iter(qn("w:instrText"))
            if node.text.strip().startswith("TOC")
        ]

    plain = _export("正文")
    assert toc_instructions(plain) == [] and "目录" not in _texts(plain)

    document = _export("正文", include_toc=True, project_overview="概述内容")
    assert toc_instructions(document) == ['TOC \\o "1-3" \\h \\z \\u']
    toc_title = _para(document, "目录")
    assert not toc_title.style.name.startswith("Heading"), "「目录」自己不应进入目录"
    texts = _texts(document)
    assert texts.index("目录") < texts.index("项目概述") < texts.index("1 技术方案")


# ---------------------------------------------------------------------------
# 配图
# ---------------------------------------------------------------------------


def _png_b64(width: int = 80, height: int = 40) -> str:
    buffer = io.BytesIO()
    Image.new("RGB", (width, height), (30, 120, 200)).save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def test_chapter_image_is_embedded():
    outline = [
        WordExportOutlineItem(
            id="1",
            title="系统架构",
            content="架构说明。",
            image_b64=_png_b64(),
            image_title="系统架构图",
        ),
        WordExportOutlineItem(id="2", title="坏图章节", content="正文仍在。", image_b64="不是 base64！"),
    ]
    document = _export(outline=outline)

    assert len(document.inline_shapes) == 1  # 解码失败的图只跳过，不影响导出
    texts = _texts(document)
    assert texts.index("架构说明。") < texts.index("系统架构图") < texts.index("2 坏图章节")
    assert _para(document, "系统架构图").alignment == WD_ALIGN_PARAGRAPH.CENTER
    _para(document, "正文仍在。")


def test_image_fits_inside_page():
    outline = [WordExportOutlineItem(id="1", title="流程", content="流程说明。", image_b64=_png_b64(40, 400))]
    document = _export(outline=outline)

    section = document.sections[0]
    shape = document.inline_shapes[0]
    assert shape.width <= section.page_width - section.left_margin - section.right_margin
    assert shape.height <= section.page_height - section.top_margin - section.bottom_margin
    assert abs(shape.height / shape.width - 10) < 0.1  # 等比缩放


# ---------------------------------------------------------------------------
# 性质：不丢字
# ---------------------------------------------------------------------------

_KITCHEN_SINK = "\r\n".join(
    [
        "**（一）总体思路**：坚持统筹规划、分步实施。",
        "-5℃ 以下需采取保温措施；+3dB 为增益上限。",
        "*注：带*号条款为实质性要求，带**号的为加分项。",
        "*整行斜体的备注*",
        "",
        "# 实施范围",
        "##建设目标",
        "#3 号楼不在本期范围内",
        "",
        "| 序号 | 设备 | 数量 | 说明 |",
        "| --- | :---: | ---: | --- |",
        "| 1 | 核心交换机 | 2 | 双机热备<br>支持堆叠 |",
        "| 2 | | | 待定 |",
        "| 3 | 防火墙 | 1 | 含入侵检测 | 多余单元格内容 |",
        "| 4 |",
        "表格后的说明文字。",
        "",
        "指标甲 | 指标乙 | 指标丙",
        "",
        "1. 准备阶段",
        "   - 现场勘察",
        "     - 机房环境",
        "   - 资料收集",
        "",
        "     收集范围包括图纸与台账",
        "2. 实施阶段",
        "+ 加号列表",
        "",
        "> 引用：以招标文件为准。",
        "> - 引用中的列表",
        ">",
        "> 引用第二段",
        "",
        "```yaml",
        "名称: 示例  # 注释也要保留",
        "  端口: 8080",
        "```",
        "```以下为配置片段",
        "参数 = 值",
        "```",
        "---",
        "- - -",
        "其中`行内代码`、**加粗*夹星号*文本**、7\\*24 小时、2*3=6、138****0000 均应保留。",
        "含控制字符\x00\x07的行，含\ttab 的行，含 &nbsp; 实体与 <br> 换行的行。",
        "日期：____年__月__日（盖章） C:\\目录\\文件",
        "~~删除线不处理~~ [链接文字](http://example.com) ![图片说明](a.png)",
        "　　全角空格缩进的段落。",
        "```",
        "没有闭合的围栏之后的文字",
    ]
)


def test_no_chinese_character_is_lost():
    outline = [WordExportOutlineItem(id="1", title="T", content=_KITCHEN_SINK)]
    document = _export(outline=outline, project_name="P")
    exported = _all_text(document)

    missing = sorted({ch for ch in _cjk(_KITCHEN_SINK) if ch not in exported})
    assert not missing, f"导出后丢失的汉字：{''.join(missing)}"
    # 更强的性质：汉字不多不少，顺序也不变
    assert _cjk(exported) == _cjk(_KITCHEN_SINK)
    # 表格之外的内容没有被误塞进表格
    assert len(document.tables) == 1 and len(document.tables[0].rows) == 5


_FUZZ_PREFIXES = [
    "", "", "", "- ", "* ", "+ ", "1. ", "12) ", "> ", ">", "# ", "## ", "###", "-", "*", "**", "***",
    "|", "| ", "    ", "  - ", "\t- ", "   1. ", "```", "```python", "~~~", "---", "- - -", "===",
    "| --- | --- |", "|---|", "|:-:|", "<br>", "&nbsp;", "\\*", "`", "``", "_", "　　", ">>",
]
_FUZZ_TOKENS = [
    "*", "**", "***", "`", "``", "|", " | ", "\\", "\\|", "\\*", "<br>", "<br/>", "&amp;", "&", "<", ">",
    " ", "  ", "\t", "#", "-", "_", "__", "~~", "1.", "2*3", "a_b", "：", "，", "。", "（", "）", "\x00", "\x0b",
]
_FUZZ_HANZI = "标书技术方案系统架构设备清单实施进度质量安全保障措施服务响应培训验收运维数据接口网络平台"


def _fuzz_document(rng: random.Random) -> str:
    lines = []
    for _ in range(rng.randint(1, 30)):
        if rng.random() < 0.12:
            lines.append("")
            continue
        parts = [rng.choice(_FUZZ_PREFIXES)]
        for _ in range(rng.randint(0, 6)):
            if rng.random() < 0.55:
                parts.append("".join(rng.choice(_FUZZ_HANZI) for _ in range(rng.randint(1, 4))))
            else:
                parts.append(rng.choice(_FUZZ_TOKENS))
        lines.append("".join(parts))
    return rng.choice(["\n", "\r\n"]).join(lines)


def test_random_markdown_never_loses_chinese_characters():
    rng = random.Random(20240917)
    for round_index in range(120):
        source = _fuzz_document(rng)
        outline = [WordExportOutlineItem(id="1", title="T", content=source)]
        document = _export(outline=outline, project_name="P")
        assert _cjk(_all_text(document)) == _cjk(source), f"第 {round_index} 轮丢字或乱序：{source!r}"
