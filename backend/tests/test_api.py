"""项目工作台 API 的端到端测试：用假模型把整条流程走一遍。"""
from __future__ import annotations

import io
import json
import time
from typing import Any, Dict, List

import docx
import pytest
from fastapi.testclient import TestClient

import app.core.jobs as jobs_module
import app.routers.projects as projects_router
from app.core.fake_llm import FakeLLM
from app.core.store import ProjectStore
from app.utils.errors import AppError
from app.writing import writer as writer_module

from .helpers import OUTLINE, REQUIREMENTS_TEXT, TENDER_TEXT, WritingResponder


def _docx_bytes(paragraphs: List[str]) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _AnalysisResponder(WritingResponder):
    """在写作应答器的基础上，再覆盖解析与目录生成的请求。"""

    def __call__(self, call):
        system = call.system
        if "招标文件分析助手" in system:
            return "项目概述：智慧园区管理平台，服务期 180 日历天。" if "提取并总结项目概述信息" in call.last_user else REQUIREMENTS_TEXT
        if "目录审核专家" in system:
            return json.dumps({"passed": True, "suggestions": []})
        if "标书编写专家" in system:
            return json.dumps({"outline": _with_fields(OUTLINE)}, ensure_ascii=False)
        return super().__call__(call)


def _with_fields(nodes, prefix=""):
    result = []
    for index, node in enumerate(nodes, start=1):
        node_id = f"{prefix}.{index}" if prefix else str(index)
        item = {"id": node_id, "title": node["title"], "description": node.get("description", "")}
        if node.get("children"):
            item["children"] = _with_fields(node["children"], node_id)
        result.append(item)
    return result


@pytest.fixture()
def api(tmp_path, monkeypatch):
    """独立的数据目录 + 假模型 + TestClient。"""
    monkeypatch.setenv("AI_BID_HOME", str(tmp_path / "home"))
    monkeypatch.setattr(writer_module, "RETRY_DELAYS", (0.0, 0.0, 0.0))
    store = ProjectStore(tmp_path / "home" / "projects")
    manager = jobs_module.JobManager(store)
    llm = FakeLLM(_AnalysisResponder())
    monkeypatch.setattr(projects_router, "store", store)
    monkeypatch.setattr(projects_router, "job_manager", manager)
    monkeypatch.setattr(projects_router, "create_llm", lambda model=None: llm)

    from app.main import app

    with TestClient(app) as client:
        client.llm = llm  # type: ignore[attr-defined]
        client.store = store  # type: ignore[attr-defined]
        yield client


def _sse_events(response) -> List[Dict[str, Any]]:
    events = []
    for block in response.text.split("\n\n"):
        for line in block.split("\n"):
            if line.startswith("data: ") and line[6:] != "[DONE]":
                events.append(json.loads(line[6:]))
    return events


def _wait_job(client, project_id: str, job_id: str) -> List[Dict[str, Any]]:
    """订阅任务事件流直到结束，返回全部事件。"""
    with client.stream("GET", f"/api/projects/{project_id}/jobs/{job_id}/events") as response:
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/event-stream")
        body = "".join(response.iter_text())
    events = []
    for block in body.split("\n\n"):
        seq = None
        for line in block.split("\n"):
            if line.startswith("id: "):
                seq = int(line[4:])
            elif line.startswith("data: "):
                events.append({"seq": seq, **json.loads(line[6:])})
    return events


def _prepare_project(client) -> Dict[str, Any]:
    state = client.post("/api/projects", json={"name": ""}).json()
    pid = state["project"]["id"]
    files = {"file": ("智慧园区招标文件.docx", _docx_bytes(TENDER_TEXT.split("\n")), "application/octet-stream")}
    state = client.post(f"/api/projects/{pid}/tender", files=files).json()
    assert state["project"]["name"] == "智慧园区招标文件"
    for kind in ("overview", "requirements"):
        response = client.post(f"/api/projects/{pid}/analyze", json={"analysis_type": kind})
        assert _sse_events(response)[-1]["type"] == "result"
    response = client.post(f"/api/projects/{pid}/outline/generate", json={"mode": "free"})
    events = _sse_events(response)
    assert events[-1]["type"] == "result", events
    return client.get(f"/api/projects/{pid}").json()


def test_full_flow_from_tender_to_word(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    assert state["project"]["tender"]["char_count"] > 50
    assert "180 日历天" in state["project"]["overview"]
    assert [r["id"] for r in state["requirements"]] == ["E1", "E2", "R1"]
    assert len(state["leaves"]) == 5 and state["chapters"] == {} and state["blueprint"] is None
    assert sum(state["budget"].values()) == 40 * 750

    api.patch(f"/api/projects/{pid}", json={"settings": {"target_pages": 4, "words_per_page": 500, "concurrency": 3}})

    # 规划
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "plan", "params": {}}).json()
    events = _wait_job(api, pid, snapshot["job"]["id"])
    assert events[-1]["type"] == "job_end" and events[-1]["job"]["status"] == "done"
    state = api.get(f"/api/projects/{pid}").json()
    assert state["blueprint"]["fact_sheet"]["facts"][0]["key"] == "服务期"
    assert state["blueprint_stale"] is False and len(state["blueprint"]["plans"]) == 5

    # 写作
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}}).json()
    events = _wait_job(api, pid, snapshot["job"]["id"])
    assert [e["seq"] for e in events] == list(range(1, len(events) + 1))
    assert events[-1]["job"]["progress"] == {"total": 5, "done": 5, "failed": 0, "chars": events[-1]["job"]["progress"]["chars"]}
    state = api.get(f"/api/projects/{pid}").json()
    assert len(state["chapters"]) == 5
    assert all(c["status"] == "done" and c["chars"] > 0 and "versions" not in c for c in state["chapters"].values())
    assert state["review"] is not None and state["review"]["stats"]["done"] == 5
    assert state["job"]["job"]["status"] == "done" and state["job"]["live"] == {}

    # 手动编辑 → 历史版本 → 恢复
    uid = state["leaves"][0]["uid"]
    original = state["chapters"][uid]["content"]
    edited = api.put(f"/api/projects/{pid}/chapters/{uid}", json={"content": "手动改写的正文。\n\n### 小节\n\n内容。"}).json()
    assert edited["source"] == "manual" and edited["version_count"] == 1 and edited["section_titles"] == ["小节"]
    detail = api.get(f"/api/projects/{pid}/chapters/{uid}").json()
    assert detail["versions"][0]["content"] == original
    restored = api.post(f"/api/projects/{pid}/chapters/{uid}/restore", json={"index": 0}).json()
    assert restored["content"] == original and restored["versions"][0]["content"].startswith("手动改写的正文")

    # 修订
    snapshot = api.post(
        f"/api/projects/{pid}/jobs", json={"kind": "revise", "params": {"uid": uid, "instruction": "补充验收方式"}}
    ).json()
    events = _wait_job(api, pid, snapshot["job"]["id"])
    assert events[-1]["job"]["status"] == "done"
    assert api.get(f"/api/projects/{pid}/chapters/{uid}").json()["source"] == "revised"

    # 审校
    report = api.post(f"/api/projects/{pid}/review/run").json()
    assert report["stats"]["leaf_total"] == 5 and len(report["coverage"]) == 3

    # 导出
    response = api.post(f"/api/projects/{pid}/export", json={})
    assert response.status_code == 200, response.text
    assert response.headers["x-missing-chapters"] == "0"
    document = docx.Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "1.1.1 总体架构" in text and "2.1.2 质量保障" in text
    assert "项目概述" not in text  # 默认不把解析出的概述放进标书


def test_analysis_result_drops_model_thinking(api, monkeypatch):
    state = api.post("/api/projects", json={"name": "p"}).json()
    pid = state["project"]["id"]
    files = {"file": ("招标文件.docx", _docx_bytes(TENDER_TEXT.splitlines()), "application/octet-stream")}
    api.post(f"/api/projects/{pid}/tender", files=files)
    original = api.llm.responder
    api.llm.responder = lambda call: (
        "<think>先看看招标文件讲了什么……</think>" + chr(10) + "项目概述：建设综合管理平台。"
        if "提取并总结项目概述信息" in call.last_user
        else original(call)
    )
    events = _sse_events(api.post(f"/api/projects/{pid}/analyze", json={"analysis_type": "overview"}))
    assert events[-1]["text"] == "项目概述：建设综合管理平台。"
    assert api.get(f"/api/projects/{pid}").json()["project"]["overview"] == "项目概述：建设综合管理平台。"


def test_export_blocks_incomplete_unless_allowed(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    response = api.post(f"/api/projects/{pid}/export", json={})
    assert response.status_code == 400 and "还有章节没有正文" in response.json()["detail"]
    response = api.post(f"/api/projects/{pid}/export", json={"allow_incomplete": True, "include_overview": True})
    assert response.status_code == 200 and response.headers["x-missing-chapters"] == "5"
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs)
    assert "项目概述" in text


def test_outline_edit_keeps_chapters_and_marks_blueprint_stale(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "plan", "params": {}}).json()
    _wait_job(api, pid, snapshot["job"]["id"])
    uid = state["leaves"][0]["uid"]
    api.put(f"/api/projects/{pid}/chapters/{uid}", json={"content": "已有正文。"})

    outline = state["project"]["outline"]
    outline[0]["title"] = "总体技术方案（修订）"
    outline.append({"title": "运维服务", "description": "", "children": [{"title": "运维体系", "description": ""}]})
    state = api.put(f"/api/projects/{pid}/outline", json={"outline": outline}).json()
    assert len(state["leaves"]) == 6 and state["blueprint_stale"] is True
    assert state["chapters"][uid]["content"] == "已有正文。"  # 改目录不丢正文
    assert state["leaves"][-1]["id"] == "3.1"

    # 只给新增章节补计划
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "plan", "params": {"only_missing": True}}).json()
    _wait_job(api, pid, snapshot["job"]["id"])
    state = api.get(f"/api/projects/{pid}").json()
    assert state["blueprint_stale"] is False and len(state["blueprint"]["plans"]) == 6

    # 重新生成目录：标题路径相同的章节沿用原来的标识，正文还在
    response = api.post(f"/api/projects/{pid}/outline/generate", json={"mode": "free"})
    assert _sse_events(response)[-1]["type"] == "result"
    state = api.get(f"/api/projects/{pid}").json()
    kept = [leaf for leaf in state["leaves"] if leaf["title"] == "进度安排"][0]
    assert kept["uid"] in state["blueprint"]["plans"]

    state = api.post(f"/api/projects/{pid}/outline/response-chapters").json()
    assert sum(1 for leaf in state["leaves"] if leaf["is_response_chapter"]) == 5
    assert api.post(f"/api/projects/{pid}/outline/response-chapters").status_code == 400


def test_blueprint_can_be_edited(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    uid = state["leaves"][0]["uid"]
    body = {
        "fact_sheet": {"project_name": "手填项目", "facts": [{"key": "工期", "value": "180 日历天", "source": "tender"}]},
        "plans": {
            uid: {"uid": "ignored", "key_points": ["要点"], "requirement_ids": ["E1", "E99"], "weight": 9, "note": "多写架构图说明"},
            "nonexistent": {"uid": "x"},
        },
    }
    state = api.put(f"/api/projects/{pid}/blueprint", json=body).json()
    plan = state["blueprint"]["plans"][uid]
    assert plan["requirement_ids"] == ["E1"] and plan["weight"] == 5.0 and plan["note"] == "多写架构图说明"
    assert "nonexistent" not in state["blueprint"]["plans"]
    assert state["blueprint"]["unassigned_requirements"] == ["E2", "R1"]
    assert state["budget"][uid] > state["budget"][state["leaves"][1]["uid"]]  # 权重影响预算
    assert api.delete(f"/api/projects/{pid}/blueprint").json()["blueprint"] is None


def test_job_conflicts_cancel_and_resume(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    api.llm.delay = 0.02
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}}).json()
    job_id = snapshot["job"]["id"]
    time.sleep(0.2)

    conflict = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}})
    assert conflict.status_code == 409
    assert api.put(f"/api/projects/{pid}/outline", json={"outline": state["project"]["outline"]}).status_code == 409

    current = api.get(f"/api/projects/{pid}/jobs/current").json()["job"]
    assert current["job"]["status"] == "running" and current["live"]
    busy_uid = next(iter(current["live"]))
    assert api.put(f"/api/projects/{pid}/chapters/{busy_uid}", json={"content": "x"}).status_code == 409
    assert api.get("/api/projects").json()["projects"][0]["busy"] is True

    cancelled = api.post(f"/api/projects/{pid}/jobs/{job_id}/cancel").json()
    assert cancelled["job"]["status"] == "cancelled" and cancelled["live"] == {}

    # 断线后从某个序号之后补发
    tail = _wait_job_after(api, pid, job_id, cancelled["last_seq"] - 2)
    assert [e["seq"] for e in tail] == [cancelled["last_seq"] - 1, cancelled["last_seq"]]
    assert tail[-1]["type"] == "job_end"

    api.llm.delay = 0
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}}).json()
    events = _wait_job(api, pid, snapshot["job"]["id"])
    assert events[-1]["job"]["status"] == "done"
    state = api.get(f"/api/projects/{pid}").json()
    assert all(c["status"] == "done" for c in state["chapters"].values()) and len(state["chapters"]) == 5


def _wait_job_after(client, project_id, job_id, after):
    with client.stream("GET", f"/api/projects/{project_id}/jobs/{job_id}/events", headers={"Last-Event-ID": str(after)}) as response:
        body = "".join(response.iter_text())
    events = []
    for block in body.split("\n\n"):
        seq = None
        for line in block.split("\n"):
            if line.startswith("id: "):
                seq = int(line[4:])
            elif line.startswith("data: "):
                events.append({"seq": seq, **json.loads(line[6:])})
    return events


def test_errors_are_reported_clearly(api, monkeypatch):
    assert api.get("/api/projects/deadbeef0000").status_code == 404
    assert api.get("/api/projects/..%2F..%2Fetc").status_code in (404, 405)
    state = api.post("/api/projects", json={"name": "空项目"}).json()
    pid = state["project"]["id"]

    assert api.post(f"/api/projects/{pid}/analyze", json={"analysis_type": "overview"}).status_code == 400
    assert api.post(f"/api/projects/{pid}/analyze", json={"analysis_type": "x"}).status_code == 400
    assert api.post(f"/api/projects/{pid}/outline/generate", json={}).status_code == 400
    assert api.post(f"/api/projects/{pid}/jobs", json={"kind": "write"}).status_code == 400  # 没有目录
    assert api.post(f"/api/projects/{pid}/jobs", json={"kind": "nope"}).status_code == 400
    assert api.get(f"/api/projects/{pid}/chapters/abcd1234").status_code == 404

    bad = {"file": ("说明.txt", b"plain text", "text/plain")}
    response = api.post(f"/api/projects/{pid}/tender", files=bad)
    assert response.status_code == 400 and "docx" in response.json()["detail"]

    # 没有配置 Key：启动任务时立刻报错，而不是进了后台才失败
    api.put(f"/api/projects/{pid}/outline", json={"outline": [{"title": "A", "children": [{"title": "A1"}]}]})

    def _no_key(model=None):
        raise AppError("请先配置OpenAI API密钥", status_code=400)

    monkeypatch.setattr(projects_router, "create_llm", _no_key)
    response = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write"})
    assert response.status_code == 400 and "API" in response.json()["detail"]
    assert api.get(f"/api/projects/{pid}/jobs/current").json()["job"] is None

    assert api.delete(f"/api/projects/{pid}").json() == {"success": True}
    assert api.delete(f"/api/projects/{pid}").status_code == 404


def test_archive_roundtrip_and_materials(api):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    files = {"file": ("旧方案.docx", _docx_bytes(["运维服务方案：7×24 小时响应。"] * 30), "application/octet-stream")}
    state = api.post(f"/api/projects/{pid}/materials", files=files).json()
    material = state["materials"][0]
    assert material["name"] == "旧方案.docx" and "text" not in material

    archive = api.get(f"/api/projects/{pid}/archive")
    assert archive.status_code == 200 and "attachment" in archive.headers["content-disposition"]
    imported = api.post("/api/projects/import", json=archive.json()).json()
    assert imported["project"]["id"] != pid and len(imported["materials"]) == 1
    assert len(api.get("/api/projects").json()["projects"]) == 2

    state = api.delete(f"/api/projects/{pid}/materials/{material['id']}").json()
    assert state["materials"] == []
    assert api.delete(f"/api/projects/{pid}/materials/{material['id']}").status_code == 404


def test_interrupted_job_is_reported_after_restart(api, monkeypatch):
    state = _prepare_project(api)
    pid = state["project"]["id"]
    snapshot = api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}}).json()
    _wait_job(api, pid, snapshot["job"]["id"])

    # 模拟“任务跑到一半服务被关掉再重启”：磁盘上的记录停在运行中，新的进程里没有这个任务
    record = api.store.load_jobs(pid, limit=1)[0]
    record.status = "running"
    api.store.save_job(record)
    fresh = jobs_module.JobManager(api.store)
    monkeypatch.setattr(projects_router, "job_manager", fresh)

    state = api.get(f"/api/projects/{pid}").json()
    assert state["job"]["job"]["status"] == "interrupted" and state["job"]["live"] == {}
    assert fresh.recover_interrupted() == 1
    assert api.get(f"/api/projects/{pid}/jobs/current").json()["job"]["job"]["status"] == "interrupted"
    # 中断不妨碍接着启动新任务
    assert api.post(f"/api/projects/{pid}/jobs", json={"kind": "write", "params": {}}).status_code == 200


def test_settings_never_return_the_key(api):
    view = api.get("/api/settings").json()
    assert view == {"api_key_set": False, "api_key_hint": "", "base_url": "", "model_name": "gpt-3.5-turbo", "demo": False}
    view = api.put("/api/settings", json={"api_key": "sk-test-1234567890abcd", "base_url": "https://example.invalid/v1", "model_name": "m1"}).json()
    assert view["api_key_set"] is True and view["api_key_hint"] == "sk-…abcd"
    assert "sk-test-1234567890abcd" not in json.dumps(view)
    # 不带 key 再保存：沿用已保存的密钥
    view = api.put("/api/settings", json={"base_url": "", "model_name": "m2"}).json()
    assert view["api_key_set"] is True and view["model_name"] == "m2" and view["base_url"] == ""
    meta = api.get("/api/meta").json()
    assert meta["demo"] is False and len(meta["personas"]) == 5


def test_oversized_upload_is_rejected_before_reading_the_body(api, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "max_file_size", 1024)
    state = api.post("/api/projects", json={"name": "p"}).json()
    pid = state["project"]["id"]
    big = {"file": ("招标文件.docx", b"x" * (2 * 1024 * 1024 + 10), "application/octet-stream")}
    for path in (f"/api/projects/{pid}/tender", f"/api/projects/{pid}/materials", "/api/document/upload"):
        response = api.post(path, files=big, headers={"Origin": "http://localhost:3000"})
        assert response.status_code == 413, path
        assert "上传上限" in response.json()["detail"]
        # 开发模式（3000 端口）下前端要能读到这条说明
        assert response.headers.get("access-control-allow-origin") == "http://localhost:3000"
    # 普通的 JSON 请求不受影响（项目归档可能比上传上限大）
    assert api.post("/api/projects", json={"name": "x" * 3000}).status_code == 200


def test_cross_site_requests_are_refused(api, monkeypatch):
    calls = []

    def _probe(body):
        calls.append(body)
        raise AssertionError("跨站请求不应该走到这里")

    monkeypatch.setattr(projects_router, "_probe_client", _probe)
    evil = {"Origin": "https://evil.example"}
    # 不带 Content-Type 的请求体：浏览器视为“简单请求”不发预检，而 FastAPI 仍按 JSON 解析
    payload = b'{"base_url": "https://evil.example/v1"}'
    for path in ("/api/settings/test", "/api/settings/models", "/api/config/save", "/api/projects"):
        response = api.post(path, content=payload, headers=evil)
        assert response.status_code == 403, path
        assert "其他网站" in response.json()["detail"]
    assert api.post("/api/settings/test", content=payload, headers={"Origin": "null"}).status_code == 403
    files = {"file": ("x.docx", b"PK", "application/octet-stream")}
    assert api.post("/api/document/upload", files=files, headers=evil).status_code == 403
    assert api.put("/api/settings", json={"base_url": "", "model_name": "m"}, headers=evil).status_code == 403
    assert calls == []

    # 读取类请求不受影响（跨站读不到响应，由 CORS 负责）
    assert api.get("/api/settings", headers=evil).status_code == 200
    # 本应用自己的页面（同源）和开发服务器照常可用
    assert api.post("/api/projects", json={"name": "a"}, headers={"Origin": "http://testserver"}).status_code == 200
    assert api.post("/api/projects", json={"name": "b"}, headers={"Origin": "http://localhost:3000"}).status_code == 200
    # 不是浏览器发的请求（没有 Origin）照常可用
    assert api.post("/api/projects", json={"name": "c"}).status_code == 200


def test_host_header_and_static_traversal_are_guarded(api):
    assert api.get("/health", headers={"Host": "evil.example.com"}).status_code == 400
    assert api.get("/health", headers={"Host": "127.0.0.1:8000"}).status_code == 200
    for path in ("/..%2Frequirements.txt", "/%2e%2e/requirements.txt", "/..%5Crequirements.txt", "/C:%5CWindows%5Cwin.ini"):
        response = api.get(path)
        assert "fastapi==" not in response.text and "[fonts]" not in response.text.lower()
