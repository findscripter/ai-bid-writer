"""经典模式（原三步向导）接口的回归测试：模型调用层改过之后，旧流程要照常可用。"""
from __future__ import annotations

import io
import json
from typing import Any, Dict, List

import docx
import pytest
from fastapi.testclient import TestClient

import app.services.analysis_service as analysis_service
import app.services.content_service as content_service
import app.services.outline_service as outline_service
from app.core.fake_llm import FakeCall, FakeLLM

OUTLINE = {
    "outline": [
        {
            "id": "1",
            "title": "总体方案",
            "description": "总体",
            "children": [
                {
                    "id": "1.1",
                    "title": "系统架构",
                    "description": "架构",
                    "children": [
                        {"id": "1.1.1", "title": "逻辑架构", "description": "分层"},
                        {"id": "1.1.2", "title": "部署架构", "description": "部署"},
                    ],
                }
            ],
        }
    ]
}


def classic_responder(call: FakeCall) -> Any:
    system = call.system
    if "招标文件分析助手" in system:
        return "项目概述：建设综合管理平台。"
    if "目录审核专家" in system:
        return json.dumps({"passed": True, "suggestions": []})
    if "标书目录重构专家" in system or "生成投标文件中技术标部分的目录结构" in system:
        return json.dumps(OUTLINE, ensure_ascii=False)
    if "负责为投标文件的技术标部分生成具体内容" in system:
        return "本章正文：系统采用分层架构，各层通过标准接口通信。" * 3
    return ""


@pytest.fixture()
def client(monkeypatch):
    llm = FakeLLM(classic_responder)
    factory = lambda model_name=None: llm  # noqa: E731
    for module in (analysis_service, content_service, outline_service):
        monkeypatch.setattr(module, "OpenAIUtil", factory)

    from app.main import app

    with TestClient(app) as test_client:
        test_client.llm = llm  # type: ignore[attr-defined]
        yield test_client


def _events(text: str) -> List[Dict[str, Any]]:
    events = []
    for block in text.split("\n\n"):
        for line in block.split("\n"):
            if line.startswith("data: ") and line[6:] != "[DONE]":
                events.append(json.loads(line[6:]))
    return events


def test_classic_upload_and_analyze(client):
    document = docx.Document()
    document.add_paragraph("第一章 采购公告：服务期 180 日历天。" * 20)
    buffer = io.BytesIO()
    document.save(buffer)
    response = client.post(
        "/api/document/upload",
        files={"file": ("招标文件.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    assert response["success"] is True and "180 日历天" in response["file_content"]

    stream = client.post("/api/document/analyze-stream", json={"file_content": "招标全文", "analysis_type": "overview"})
    events = _events(stream.text)
    assert "".join(e.get("chunk", "") for e in events) == "项目概述：建设综合管理平台。"
    assert stream.text.rstrip().endswith("data: [DONE]")


def test_classic_outline_stream(client):
    stream = client.post("/api/outline/generate-stream", json={"overview": "概述", "requirements": "评分要求", "mode": "free"})
    events = _events(stream.text)
    result = [e for e in events if e.get("type") == "result"]
    assert result and result[0]["outline"]["outline"][0]["children"][0]["children"][1]["id"] == "1.1.2"
    assert any(e.get("type") == "progress" for e in events)


def test_classic_chapter_stream(client):
    chapter = {"id": "1.1.1", "title": "逻辑架构", "description": "分层"}
    stream = client.post(
        "/api/content/generate-chapter-stream",
        json={"chapter": chapter, "project_overview": "概述", "target_words": 50},
    )
    text = "".join(e.get("chunk", "") for e in _events(stream.text))
    assert text.startswith("本章正文") and not any(e.get("error") for e in _events(stream.text))


def test_classic_export_word(client):
    outline = json.loads(json.dumps(OUTLINE))["outline"]
    for leaf in outline[0]["children"][0]["children"]:
        leaf["content"] = "正文段落。\n\n| 指标 | 要求 |\n| --- | --- |\n| 响应时间 | ≤3 秒 |"
    body = {"project_name": "示例项目", "outline": outline}
    response = client.post("/api/document/export-word", json=body)
    assert response.status_code == 200
    document = docx.Document(io.BytesIO(response.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "1.1.1 逻辑架构" in text

    outline[0]["children"][0]["children"][1]["content"] = ""
    blocked = client.post("/api/document/export-word", json={"project_name": "示例项目", "outline": outline})
    assert blocked.status_code == 400 and "1.1.2" in blocked.json()["detail"]


def test_classic_routes_need_a_key_when_not_patched(monkeypatch):
    """不注入假模型时，没有配置密钥要给出清楚的提示，而不是 500。"""
    from app.main import app

    with TestClient(app) as test_client:
        response = test_client.post("/api/outline/generate", json={"overview": "概述", "requirements": "评分要求"})
        assert response.status_code == 400 and "API" in response.json()["detail"]
