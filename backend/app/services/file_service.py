"""文件处理服务：识别上传文档的真实类型，并在工作线程里抽取文本。"""
from __future__ import annotations

import asyncio
import gc
import io
import logging
import os
import re
import tempfile
import time
import unicodedata
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from typing import Any, BinaryIO, Callable, Dict, Iterable, List, Optional

import docx
import PyPDF2
from docx.table import Table as DocxTable
from fastapi import UploadFile
from PyPDF2.errors import DependencyError

from ..config import settings
from ..utils.errors import AppError

logger = logging.getLogger(__name__)

# 抽取效果更好的解析库；缺了任何一个就退回 python-docx / PyPDF2
try:
    import fitz  # PyMuPDF
    import pdfplumber
    from docx2python import docx2python
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    HAS_ADVANCED_LIBS = True
except ImportError as e:
    HAS_ADVANCED_LIBS = False
    logger.warning("高级文档处理库未安装: %s", e)


@dataclass
class ExtractedDocument:
    """一次文档抽取的结果。

    text 一个字都没抽到时是空字符串；char_count 恒等于 len(text)；
    warnings 是可以直接展示给使用者的提示（疑似扫描件、内容过少等），没有问题时为空列表。
    """

    filename: str
    text: str
    char_count: int
    warnings: List[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# 常量
# ---------------------------------------------------------------------------

_SUPPORTED_EXTENSIONS = (".docx", ".pdf")
# 分块读取上传内容的块大小：既是超限判断的粒度，也是内存里同时持有的最大数据量
_CHUNK_SIZE = 256 * 1024
_SNIFF_BYTES = 2048  # 判断文件类型至少要看到的文件头长度
_TEMP_PREFIX = "ai_bid_upload_"
# 文件名里不该出现的字符：控制字符、零宽字符与文字方向控制符（C*），行、段分隔符（Zl / Zp）
_HIDDEN_CATEGORIES = ("Cc", "Cf", "Cs", "Co", "Cn", "Zl", "Zp")

_ZIP_MAGIC = b"PK\x03\x04"
_PDF_MAGIC = b"%PDF-"
_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# OLE2 容器里装的是什么，要看目录里的流名（UTF-16LE）：
# 老式 .doc 有 WordDocument，加了打开密码的新格式文档有 EncryptedPackage
_OLE2_WORD_MARK = "WordDocument".encode("utf-16-le")
_OLE2_ENCRYPTED_MARK = "EncryptedPackage".encode("utf-16-le")

# docx 是压缩包，解压后超过这个体积的文档解析时会耗尽内存，直接拒绝
_MAX_DOCX_UNPACKED_BYTES = 300 * 1024 * 1024

_SCANNED_AVG_CHARS_PER_PAGE = 50  # PDF 页均有效字符低于它，基本可以断定是扫描件
_SPARSE_PAGE_CHARS = 20  # 单页有效字符低于它，视为“几乎没有文字的页”
_SHORT_TEXT_CHARS = 200  # 全文有效字符低于它，视为内容过少

# docx2python 在图片位置留下的占位标记，形如 ----media/image1.png----
_IMAGE_MARK_RE = re.compile(r"----.*?(?:image|img|media).*?----", re.IGNORECASE)

# 解析全部排进同一个工作线程：PyMuPDF 不支持多线程并发调用，串行执行同时给内存占用封了顶。
_PARSE_EXECUTOR = ThreadPoolExecutor(max_workers=1, thread_name_prefix="doc-parse")


# ---------------------------------------------------------------------------
# 面向使用者的文案
# ---------------------------------------------------------------------------

_SUPPORTED_HINT = (
    "目前只支持 Word 文档（.docx）和 PDF（.pdf）；"
    "老式 .doc 文件请先用 Word 或 WPS 另存为 .docx 再上传。"
)
_MSG_LEGACY_DOC_EXTENSION = (
    "暂不支持老式 Word 格式（.doc）。请用 Word 或 WPS 打开这个文件，"
    "另存为 .docx 后再上传；目前只支持 .docx 和 .pdf。"
)
_MSG_LEGACY_DOC_CONTENT = (
    "这个文件的内容实际上是老式 Word 格式（.doc），无法直接读取（只改扩展名不行）。"
    "请用 Word 或 WPS 打开它，另存为 .docx 后再上传。"
)
_MSG_LEGACY_OFFICE_CONTENT = (
    f"这个文件实际上是老式 Office 格式的文件，不是 Word 文档（.docx）或 PDF。{_SUPPORTED_HINT}"
)
_MSG_ENCRYPTED_OFFICE = (
    "这个文档设置了打开密码（已加密），无法读取内容。"
    "请用 Word 或 WPS 打开它，取消密码保护并另存为 .docx 后再上传。"
)
_MSG_EMPTY_FILE = "这个文件是空的（0 字节），请确认文件完整后重新上传。"
_MSG_NOT_DOCX = (
    "这个文件不是有效的 Word 文档（.docx）：文件可能已损坏，或者是把其他格式直接改了扩展名。"
    "请用 Word 或 WPS 打开确认，另存为 .docx 后再上传。"
)
_MSG_NOT_PDF = (
    "这个文件不是有效的 PDF：文件可能已损坏，或者是把其他格式直接改了扩展名。"
    "请用 PDF 阅读器打开确认后重新上传。"
)
_MSG_DOCX_BROKEN = (
    "这个 Word 文档已损坏或格式不规范，无法读取。"
    "请用 Word 或 WPS 打开它，另存为新的 .docx 后再上传。"
)
_MSG_DOCX_TOO_BIG_UNPACKED = "这个 Word 文档解压后的内容过于庞大，无法处理。请拆分文档后再上传。"
_MSG_PDF_BROKEN = (
    "这个 PDF 文件已损坏或格式不规范，无法读取。"
    "请用 PDF 阅读器打开确认，可以尝试“打印为 PDF”重新生成后再上传。"
)
_MSG_PDF_ENCRYPTED = (
    "这个 PDF 设置了打开密码（已加密），无法读取内容。"
    "请先用 PDF 阅读器输入密码打开，另存或“打印为 PDF”得到不带密码的版本后再上传。"
)
_OCR_ADVICE = "建议换用可以复制文字的版本，或先用 OCR 工具识别成文字后再上传。"


def _format_size(num_bytes: int) -> str:
    """把字节数写成使用者习惯的 MB / KB。"""
    if num_bytes >= 1024 * 1024:
        return f"{num_bytes / (1024 * 1024):.1f}".rstrip("0").rstrip(".") + "MB"
    return f"{max(num_bytes // 1024, 1)}KB"


def _too_large_message(limit: int) -> str:
    return (
        f"文件太大：超过了 {_format_size(limit)} 的上限。"
        "请压缩文件（比如压缩或删除其中的大图片）或拆分后再上传。"
    )


def _format_pages(pages: List[int], max_ranges: int = 6) -> str:
    """把页码列表压成“第 3、7-9、12 页”这样的写法，太长时截断。"""
    ranges: List[str] = []
    start = prev = pages[0]
    for page in pages[1:] + [None]:  # 末尾的 None 用来收尾最后一段
        if page is not None and page == prev + 1:
            prev = page
            continue
        ranges.append(str(start) if start == prev else f"{start}-{prev}")
        if page is not None:
            start = prev = page
    more = "等" if len(ranges) > max_ranges else ""
    return f"第 {'、'.join(ranges[:max_ranges])} 页{more}"


# ---------------------------------------------------------------------------
# 解析结果（工作线程产出）
# ---------------------------------------------------------------------------


def _content_len(text: str) -> int:
    """有实际内容的字符数：不计空白，也不计 docx2python 留下的图片占位标记。"""
    if "----" in text:
        text = _IMAGE_MARK_RE.sub("", text)
    return len("".join(text.split()))


def _one_line(text: str) -> str:
    """把单元格里的多行文字并成一行，免得把“一行一条记录”的表格结构打散。"""
    return " ".join(part.strip() for part in text.splitlines() if part.strip())


@dataclass
class _Parsed:
    """解析线程的产出：文本片段，加上生成提示所需的统计。"""

    kind: str  # "docx" 或 "pdf"
    parts: List[str] = field(default_factory=list)
    # 正文的有效字符数；页码、表格这些我们自己加的结构标记不算在内
    content_chars: int = 0
    page_chars: List[int] = field(default_factory=list)  # 仅 PDF：每页的有效字符数
    image_count: int = 0  # 仅 docx：内嵌图片数
    images: Dict[str, bytes] = field(default_factory=dict)  # 仅在开启图片外传时填充
    degraded: bool = False  # 主解析器失败，内容是备用解析器读出来的
    text: str = ""

    def add_paragraph(self, text: str) -> None:
        text = text.strip()
        if text:
            self.parts.append(text)
            self.content_chars += _content_len(text)

    def add_table(self, number: int, rows: List[List[str]]) -> None:
        self.parts.append(f"\n[表格 {number}]")
        for cells in rows:
            self.parts.append(" | ".join(cells))
            self.content_chars += sum(_content_len(cell) for cell in cells)
        self.parts.append("[表格结束]\n")

    def add_pdf_page(self, number: int, text: str, tables: Iterable[Any]) -> None:
        self.parts.append(f"\n--- 第 {number} 页 ---\n")
        # 表格里的字 extract_text 已经给过一遍，这里只按正文统计
        chars = _content_len(text)
        self.page_chars.append(chars)
        self.content_chars += chars
        if text:
            self.parts.append(text)
        for table_number, table in enumerate(tables, 1):
            self.parts.append(f"\n[表格 {table_number}]")
            for row in table or []:
                if row:
                    self.parts.append(
                        " | ".join(str(cell) if cell else "" for cell in row)
                    )
            self.parts.append("[表格结束]\n")

    def finish(self) -> "_Parsed":
        # 一个字都没抽到时不返回只有页码标记的空壳，调用方用 `not text` 就能判断
        self.text = "\n".join(self.parts).strip() if self.content_chars else ""
        self.parts = []
        return self


# ---------------------------------------------------------------------------
# 类型识别与落盘（工作线程内执行）
# ---------------------------------------------------------------------------


def _clean_filename(raw: Optional[str]) -> str:
    """只保留文件名本身：去掉客户端带来的路径，以及控制字符、零宽字符和换行。"""
    name = (raw or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(
        ch for ch in name if unicodedata.category(ch) not in _HIDDEN_CATEGORIES
    ).strip()
    return name or "未命名文件"


def _check_extension(extension: str) -> None:
    """扩展名这一关：明显不支持的类型不必读内容就能拒绝。"""
    if not extension or extension in _SUPPORTED_EXTENSIONS:
        return  # 没有扩展名时留给文件头判断
    if extension == ".doc":
        raise AppError(_MSG_LEGACY_DOC_EXTENSION, 400)
    raise AppError(f"不支持的文件类型（{extension}）。{_SUPPORTED_HINT}", 400)


def _rewind(source: BinaryIO) -> None:
    try:
        source.seek(0)
    except (AttributeError, OSError, ValueError):
        pass  # 不能回退的流就从当前位置读


def _read_head(source: BinaryIO) -> bytes:
    head = b""
    while len(head) < _SNIFF_BYTES:
        chunk = source.read(_CHUNK_SIZE)
        if not chunk:
            break
        head += chunk
    return head


def _ole2_message(source: BinaryIO, head: bytes, limit: int) -> str:
    """文件头是 OLE2 时，接着往下找流名，判断它是老式 .doc、加密文档还是别的老式 Office 文件。"""
    overlap = len(_OLE2_ENCRYPTED_MARK)
    total = len(head)
    window = head
    encrypted = False
    while True:
        if _OLE2_WORD_MARK in window:
            return _MSG_LEGACY_DOC_CONTENT
        encrypted = encrypted or _OLE2_ENCRYPTED_MARK in window
        chunk = source.read(_CHUNK_SIZE) if total <= limit else b""
        if not chunk:
            break
        total += len(chunk)
        window = window[-overlap:] + chunk  # 带上一块的结尾，流名跨块也能找到
    return _MSG_ENCRYPTED_OFFICE if encrypted else _MSG_LEGACY_OFFICE_CONTENT


def _resolve_kind(source: BinaryIO, extension: str, head: bytes, limit: int) -> str:
    """结合扩展名和文件头确定按哪种格式解析；对不上时给出能看懂的原因。"""
    if head.startswith(_ZIP_MAGIC):
        actual = "docx"  # 是不是真的 Word 文档，解析前还会再看压缩包里的内容
    elif _PDF_MAGIC in head[:1024]:
        actual = "pdf"  # PDF 规范允许文件头前面有少量杂字节
    else:
        actual = ""

    expected = extension.lstrip(".")
    if actual and expected in ("", actual):
        return actual

    if head.startswith(_OLE2_MAGIC):
        raise AppError(_ole2_message(source, head, limit), 400)
    if expected == "docx" and actual == "pdf":
        raise AppError(
            "文件扩展名是 .docx，但内容实际上是 PDF。请把文件扩展名改成 .pdf 后重新上传。",
            400,
        )
    if expected == "pdf" and actual == "docx":
        if b"word/" in head:
            raise AppError(
                "文件扩展名是 .pdf，但内容实际上是 Word 文档。"
                "请把文件扩展名改成 .docx 后重新上传。",
                400,
            )
        raise AppError(_MSG_NOT_PDF, 400)
    if expected == "docx":
        raise AppError(_MSG_NOT_DOCX, 400)
    if expected == "pdf":
        raise AppError(_MSG_NOT_PDF, 400)
    raise AppError(f"无法识别这个文件的类型。{_SUPPORTED_HINT}", 400)


def _spool_to_temp(source: BinaryIO, head: bytes, suffix: str, limit: int) -> str:
    """把上传内容分块写进系统临时目录；一旦超过大小上限立刻中止，不再往下读。"""
    out = tempfile.NamedTemporaryFile(prefix=_TEMP_PREFIX, suffix=suffix, delete=False)
    path = out.name
    try:
        with out:
            total = len(head)
            out.write(head)
            while True:
                chunk = source.read(_CHUNK_SIZE)
                if not chunk:
                    break
                total += len(chunk)
                if total > limit:
                    raise AppError(_too_large_message(limit), 400)
                out.write(chunk)
    except BaseException:
        FileService._safe_file_cleanup(path)
        raise
    return path


def _ingest_upload(
    source: BinaryIO, extension: str, limit: int, want_images: bool
) -> _Parsed:
    """（工作线程）读入上传内容 → 判断类型 → 落到临时文件 → 解析 → 删除临时文件。

    临时文件从创建到删除都在这一个函数里，请求中途被取消也不会留下文件。
    """
    started = time.perf_counter()
    _rewind(source)
    try:
        head = _read_head(source)
        if not head:
            raise AppError(_MSG_EMPTY_FILE, 400)
        if len(head) > limit:
            raise AppError(_too_large_message(limit), 400)
        kind = _resolve_kind(source, extension, head, limit)
        path = _spool_to_temp(source, head, "." + kind, limit)
        try:
            parsed = _parse_path(path, kind, want_images)
        finally:
            FileService._safe_file_cleanup(path)
    finally:
        _rewind(source)  # 调用方之后还想再读原文件（比如存档）时不必自己复位
    logger.info(
        "文档解析完成: 类型=%s 有效字符=%s 页数=%s 耗时=%.2fs",
        parsed.kind,
        parsed.content_chars,
        len(parsed.page_chars) or "-",
        time.perf_counter() - started,
    )
    return parsed


def _parse_path(path: str, kind: str, want_images: bool) -> _Parsed:
    """（工作线程）按类型解析本地文件。"""
    parsed = _parse_pdf(path) if kind == "pdf" else _parse_docx(path, want_images)
    return parsed.finish()


# 下面的解析函数都自己打开文件、把文件对象（或读出的数据）交给解析库，而不是把路径交出去：
# 有的库打开失败时，异常对象会一直占着它开的文件句柄，Windows 上临时文件就删不掉了。


def _attempt(label: str, parser: Callable[..., _Parsed], *args: Any) -> Optional[_Parsed]:
    """跑一个解析器；失败时记下原因并返回 None，由调用方决定换备用方案还是报错。"""
    try:
        return parser(*args)
    except AppError:
        raise
    except Exception:
        logger.warning("%s 解析失败", label, exc_info=True)
        return None


# ---------------------------------------------------------------------------
# Word 文档
# ---------------------------------------------------------------------------


def _check_docx_package(path: str) -> int:
    """确认压缩包确实是一份 Word 文档，返回内嵌图片数。"""
    try:
        with open(path, "rb") as stream, zipfile.ZipFile(stream) as package:
            members = package.infolist()
    except zipfile.BadZipFile as exc:
        raise AppError(_MSG_DOCX_BROKEN, 400) from exc

    names = [member.filename for member in members]
    if not any(name.startswith("word/") for name in names):
        if any(name.startswith("xl/") for name in names):
            looks_like = "（看起来是 Excel 表格）"
        elif any(name.startswith("ppt/") for name in names):
            looks_like = "（看起来是 PowerPoint 演示文稿）"
        else:
            looks_like = "（看起来是压缩包或其他格式）"
        raise AppError(f"这个文件不是 Word 文档{looks_like}。{_SUPPORTED_HINT}", 400)
    if sum(member.file_size for member in members) > _MAX_DOCX_UNPACKED_BYTES:
        raise AppError(_MSG_DOCX_TOO_BIG_UNPACKED, 400)
    return sum(1 for name in names if name.startswith("word/media/"))


def _is_table_block(block: List[Any]) -> bool:
    """docx2python 把正文段落也包成 1×1 的“表”，只有 lineage 里带 tbl 的才是真表格。"""
    for row in block:
        for cell in row:
            for par in cell:
                return par.lineage[1] == "tbl"
    return False


def _parse_docx_docx2python(path: str, want_images: bool) -> _Parsed:
    """用 docx2python 提取：能还原自动编号，也覆盖页眉页脚、脚注和文本框。"""
    parsed = _Parsed(kind="docx")
    table_number = 0
    with open(path, "rb") as stream, docx2python(stream) as content:
        for block in content.document_pars:
            if not _is_table_block(block):
                for row in block:
                    for cell in row:
                        for par in cell:
                            parsed.add_paragraph("".join(par.run_strings))
                continue
            rows = []
            for row in block:
                cells = [
                    _one_line("\n".join("".join(par.run_strings) for par in cell))
                    for cell in row
                ]
                if any(cells):
                    rows.append(cells)
            if rows:
                table_number += 1
                parsed.add_table(table_number, rows)
        if want_images:
            try:
                parsed.images = dict(content.images)
            except Exception:
                logger.warning("读取 Word 文档内嵌图片失败，本次不做图片外传", exc_info=True)
    return parsed


def _parse_docx_python_docx(path: str) -> _Parsed:
    """用 python-docx 提取（备用）：按文档顺序输出段落和表格。"""
    parsed = _Parsed(kind="docx")
    table_number = 0
    with open(path, "rb") as stream:
        document = docx.Document(stream)
    for block in document.iter_inner_content():
        if not isinstance(block, DocxTable):
            parsed.add_paragraph(block.text)
            continue
        rows = []
        for row in block.rows:
            cells = [_one_line(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if rows:
            table_number += 1
            parsed.add_table(table_number, rows)
    return parsed


def _parse_docx(path: str, want_images: bool = False) -> _Parsed:
    image_count = _check_docx_package(path)
    parsed: Optional[_Parsed] = None
    if HAS_ADVANCED_LIBS:
        parsed = _attempt("docx2python", _parse_docx_docx2python, path, want_images)
    if parsed is None:
        parsed = _attempt("python-docx", _parse_docx_python_docx, path)
    if parsed is None:
        raise AppError(_MSG_DOCX_BROKEN, 400)
    parsed.image_count = image_count
    return parsed


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _is_password_error(exc: BaseException) -> bool:
    """pdfplumber 把 pdfminer 的异常包了一层，“需要密码”要到里面去找。"""
    inner = exc.args[0] if exc.args else None
    return isinstance(exc, PDFPasswordIncorrect) or isinstance(inner, PDFPasswordIncorrect)


def _parse_pdf_pdfplumber(path: str) -> _Parsed:
    """用 pdfplumber 提取文本和表格。"""
    parsed = _Parsed(kind="pdf")
    with open(path, "rb") as stream:
        try:
            pdf = pdfplumber.open(stream)
        except Exception as exc:
            if _is_password_error(exc):
                raise AppError(_MSG_PDF_ENCRYPTED, 400) from exc
            raise
        with pdf:
            for number, page in enumerate(pdf.pages, 1):
                try:
                    text = page.extract_text() or ""
                    tables = page.extract_tables()
                finally:
                    page.close()  # pdfplumber 会缓存每页的版面对象，不逐页释放大文件会吃满内存
                parsed.add_pdf_page(number, text, tables)
    return parsed


def _parse_pdf_pymupdf(path: str) -> _Parsed:
    """用 PyMuPDF 提取文本和表格（备用，对结构有损坏的 PDF 更宽容）。"""
    parsed = _Parsed(kind="pdf")
    with open(path, "rb") as stream:
        data = stream.read()
    # MuPDF 修复受损文件时会往控制台打印成串的错误，先关掉，结束后收进日志
    show_errors = fitz.TOOLS.mupdf_display_errors()
    fitz.TOOLS.mupdf_display_errors(False)
    try:
        document = fitz.open(stream=data, filetype="pdf")
        try:
            if document.needs_pass:
                raise AppError(_MSG_PDF_ENCRYPTED, 400)
            for index in range(document.page_count):
                page = document[index]
                text = page.get_text() or ""
                tables: List[Any] = []
                try:
                    tables = [table.extract() for table in page.find_tables().tables]
                except Exception:
                    logger.debug("PDF 第 %s 页表格提取失败，已跳过", index + 1, exc_info=True)
                parsed.add_pdf_page(index + 1, text, tables)
        finally:
            document.close()
    finally:
        fitz.TOOLS.mupdf_display_errors(show_errors)
        notes = fitz.TOOLS.mupdf_warnings()
        if notes:
            logger.debug("MuPDF 提示: %s", notes)
    return parsed


def _parse_pdf_pypdf2(path: str) -> _Parsed:
    """用 PyPDF2 提取文本（高级库都没装时的兜底，不含表格）。"""
    parsed = _Parsed(kind="pdf")
    with open(path, "rb") as stream:
        try:
            reader = PyPDF2.PdfReader(stream)
            unlocked = not reader.is_encrypted or bool(reader.decrypt(""))
        except DependencyError:
            unlocked = False  # 加密了，而且环境里没有解密要用的库
        if not unlocked:
            raise AppError(_MSG_PDF_ENCRYPTED, 400)
        for number, page in enumerate(reader.pages, 1):
            parsed.add_pdf_page(number, page.extract_text() or "", [])
    return parsed


def _parse_pdf(path: str) -> _Parsed:
    if HAS_ADVANCED_LIBS:
        parsed = _attempt("pdfplumber", _parse_pdf_pdfplumber, path)
        if parsed is None:
            parsed = _attempt("PyMuPDF", _parse_pdf_pymupdf, path)
            # 主解析器读不了、备用解析器又一个字都没读出来，这不是扫描件而是文件坏了
            if parsed is not None and not parsed.content_chars:
                parsed = None
            if parsed is not None:
                parsed.degraded = True
    else:
        parsed = _attempt("PyPDF2", _parse_pdf_pypdf2, path)
    if parsed is None or not parsed.page_chars:
        raise AppError(_MSG_PDF_BROKEN, 400)
    return parsed


# ---------------------------------------------------------------------------
# 提示
# ---------------------------------------------------------------------------


def _build_warnings(parsed: _Parsed) -> List[str]:
    """根据抽取统计生成给使用者看的提示。"""
    warnings: List[str] = []
    chars = parsed.content_chars

    if parsed.kind == "pdf":
        pages = len(parsed.page_chars)
        sparse = [
            number
            for number, count in enumerate(parsed.page_chars, 1)
            if count < _SPARSE_PAGE_CHARS
        ]
        average = chars / pages
        mostly_sparse = len(sparse) >= 3 and len(sparse) * 2 >= pages
        if chars == 0:
            warnings.append(
                f"这个 PDF 共 {pages} 页，但没有提取到任何文字（0 个字符），"
                f"很可能是扫描件或图片版 PDF。{_OCR_ADVICE}"
            )
        elif average < _SCANNED_AVG_CHARS_PER_PAGE:
            per_page = f"约 {average:.0f} 个字符" if average >= 1 else "不到 1 个字符"
            warnings.append(
                f"这个 PDF 共 {pages} 页，只提取到 {chars} 个字符（平均每页{per_page}），"
                f"很可能是扫描件或图片版 PDF，后续的解析结果会不可靠。{_OCR_ADVICE}"
            )
        elif mostly_sparse:
            warnings.append(
                f"这个 PDF 共 {pages} 页，其中 {len(sparse)} 页几乎没有可提取的文字，"
                f"全文只提取到 {chars} 个字符，很可能大部分是扫描页或图片，"
                f"后续的解析结果会不完整。{_OCR_ADVICE}"
            )
        else:
            if len(sparse) >= 3 and len(sparse) * 5 >= pages:
                warnings.append(
                    f"这个 PDF 共 {pages} 页，其中 {len(sparse)} 页（{_format_pages(sparse)}）"
                    "几乎没有可提取的文字，可能是扫描页或图片页，这些页的内容不会被解析。"
                    "如果里面有重要内容，建议换用可以复制文字的版本，或先对这些页做 OCR。"
                )
            if chars < _SHORT_TEXT_CHARS:
                warnings.append(
                    f"只提取到 {chars} 个字符，内容非常少，请确认上传的是完整的文件。"
                )
        if parsed.degraded:
            warnings.append(
                "这个 PDF 的文件结构有损坏或不规范，已尽量读取，内容可能不完整。"
                "建议用 PDF 阅读器“打印为 PDF”重新生成后再上传。"
            )
        return warnings

    if chars == 0:
        if parsed.image_count:
            warnings.append(
                "没有从这个 Word 文档中提取到任何文字（0 个字符）。"
                f"文档里有 {parsed.image_count} 张图片，内容可能是以图片形式粘贴的；{_OCR_ADVICE}"
            )
        else:
            warnings.append(
                "没有从这个 Word 文档中提取到任何文字（0 个字符），请确认上传的不是空白文档。"
            )
    elif chars < _SHORT_TEXT_CHARS:
        message = f"只提取到 {chars} 个字符，内容非常少，请确认上传的是完整的文件。"
        if parsed.image_count:
            message += (
                f"文档里有 {parsed.image_count} 张图片，图片中的文字无法提取；"
                f"如果正文在图片里，{_OCR_ADVICE}"
            )
        warnings.append(message)
    return warnings


# ---------------------------------------------------------------------------
# 调度
# ---------------------------------------------------------------------------


async def _run_in_parser_thread(func: Callable[..., Any], *args: Any) -> Any:
    """把阻塞的解析放到专用工作线程，事件循环照常响应其他请求。"""
    return await asyncio.get_running_loop().run_in_executor(_PARSE_EXECUTOR, func, *args)


def _image_extension(name: str) -> str:
    suffix = os.path.splitext(name)[1].lower().lstrip(".")
    if suffix in ("png", "gif", "bmp"):
        return suffix
    return "jpg"


class FileService:
    """文件处理服务"""

    ALLOWED_DOCUMENT_TYPES = {
        "application/pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    SUPPORTED_EXTENSIONS = _SUPPORTED_EXTENSIONS

    # 可选：把招标文件里的插图传到外部图床。默认关闭，避免把标书图片发到第三方。
    IMAGE_UPLOAD_URL = os.environ.get("IMAGE_UPLOAD_URL", "").strip()
    IMAGE_UPLOAD_TIMEOUT = 30  # 超时时间（秒）

    @staticmethod
    def is_supported_document(content_type: str | None) -> bool:
        """按 Content-Type 粗筛上传类型。

        浏览器给的 Content-Type 并不可靠（没装 Office 的机器上 .docx 可能是空值），
        真正的类型判断在 extract_document 里按扩展名和文件头做，新代码不必再调用这里。
        """
        return bool(content_type and content_type in FileService.ALLOWED_DOCUMENT_TYPES)

    @staticmethod
    async def upload_image_to_server(image_data: bytes, filename: str) -> Optional[str]:
        """上传图片到外部服务器。未配置 IMAGE_UPLOAD_URL 时直接跳过。"""
        if not FileService.IMAGE_UPLOAD_URL:
            return None
        try:
            import aiohttp  # 只有开启图片外传才需要

            # 准备multipart/form-data格式的数据
            form_data = aiohttp.FormData()
            form_data.add_field(
                "file",
                io.BytesIO(image_data),
                filename=filename,
                content_type="image/jpeg",
            )

            timeout = aiohttp.ClientTimeout(total=FileService.IMAGE_UPLOAD_TIMEOUT)
            async with aiohttp.ClientSession(timeout=timeout) as session:
                async with session.post(
                    FileService.IMAGE_UPLOAD_URL, data=form_data
                ) as response:
                    if response.status == 200:
                        result = await response.json()
                        # 根据实际API返回格式获取图片URL
                        return result.get("file_url")
                    else:
                        logger.warning("图片上传失败，状态码: %s", response.status)
                        return None
        except Exception as e:
            logger.warning("图片上传异常: %s", e)
            return None

    @staticmethod
    async def _link_uploaded_images(text: str, images: Dict[str, bytes]) -> str:
        """把正文里的图片占位标记换成已上传图片的引用（仅在配置了 IMAGE_UPLOAD_URL 时调用）。"""
        references: List[str] = []
        labels: Dict[str, Optional[str]] = {}  # 图片名 → 引用标记；同一张图只传一次，失败记 None
        pieces: List[str] = []
        cursor = 0
        for match in _IMAGE_MARK_RE.finditer(text):
            name = os.path.basename(match.group()[4:-4].strip())
            if name not in images:
                continue
            if name not in labels:
                number = len(references) + 1
                url = await FileService.upload_image_to_server(
                    images[name], f"docx_img{number}.{_image_extension(name)}"
                )
                labels[name] = f"[图片{number}]" if url else None
                if url:
                    references.append(f"[图片{number}]: {url}")
            label = labels[name]
            if label is None:
                continue
            pieces.append(text[cursor:match.start()])
            pieces.append(label)
            cursor = match.end()
        if not references:
            return text
        pieces.append(text[cursor:])
        return "".join(pieces) + "\n\n\n--- 图片引用 ---\n" + "\n".join(references)

    @staticmethod
    async def _finish_text(parsed: _Parsed) -> str:
        if parsed.text and parsed.images and FileService.IMAGE_UPLOAD_URL:
            return await FileService._link_uploaded_images(parsed.text, parsed.images)
        return parsed.text

    @staticmethod
    def _safe_file_cleanup(file_path: str, max_retries: int = 3) -> bool:
        """删除临时文件。Windows 上文件可能被杀毒软件等短暂占用，删不掉时稍等再试。

        重试之间会 sleep，只应在工作线程里调用。
        """
        for attempt in range(max_retries):
            try:
                os.remove(file_path)
                return True
            except FileNotFoundError:
                return True
            except OSError as exc:
                if attempt == max_retries - 1:
                    logger.warning("无法删除临时文件 %s: %s", file_path, exc)
                    return False
                gc.collect()  # 解析库若漏关了句柄，回收之后才删得掉
                time.sleep(0.2 * (attempt + 1))
        return False

    @staticmethod
    async def extract_text_from_pdf(file_path: str) -> str:
        """从本地 PDF 文件提取文本（含表格），解析在工作线程里执行。"""
        parsed = await _run_in_parser_thread(_parse_path, file_path, "pdf", False)
        return parsed.text

    @staticmethod
    async def extract_text_from_docx(file_path: str) -> str:
        """从本地 Word 文档提取文本（含表格），解析在工作线程里执行。"""
        want_images = bool(FileService.IMAGE_UPLOAD_URL)
        parsed = await _run_in_parser_thread(_parse_path, file_path, "docx", want_images)
        return await FileService._finish_text(parsed)

    @staticmethod
    async def extract_document(file: UploadFile) -> ExtractedDocument:
        """抽取上传文档的文本，并给出需要提醒使用者的问题。

        - 类型按扩展名加文件头判断，不看 Content-Type；
        - 内容分块读入，超过 settings.max_file_size 立刻中止；
        - 解析在工作线程里执行，临时文件放系统临时目录并在结束时删除；
        - 使用者能自己解决的问题（类型不支持、文件过大、损坏、加密）抛 AppError(400)，
          文案可以直接展示；疑似扫描件、内容过少等放在返回值的 warnings 里。

        返回后 file 的读取位置回到开头，调用方可以接着读原文件。
        """
        filename = _clean_filename(file.filename)
        extension = os.path.splitext(filename)[1].lower()
        _check_extension(extension)

        limit = settings.max_file_size
        size = getattr(file, "size", None)
        if size is not None and size > limit:
            raise AppError(_too_large_message(limit), 400)

        want_images = bool(FileService.IMAGE_UPLOAD_URL)
        parsed = await _run_in_parser_thread(
            _ingest_upload, file.file, extension, limit, want_images
        )
        text = await FileService._finish_text(parsed)
        return ExtractedDocument(
            filename=filename,
            text=text,
            char_count=len(text),
            warnings=_build_warnings(parsed),
        )

    @staticmethod
    async def process_uploaded_file(file: UploadFile) -> str:
        """处理上传的文件并返回提取出的文本（旧入口，新代码请用 extract_document）。

        这里的调用方拿不到 warnings，所以一个字都没抽到时直接抛 AppError，其余提示只记日志。
        """
        document = await FileService.extract_document(file)
        if not document.text:
            raise AppError(
                document.warnings[0] if document.warnings else "没有从文件中提取到任何文字。",
                400,
            )
        for warning in document.warnings:
            logger.warning("文档 %s 的抽取提示: %s", document.filename, warning)
        return document.text
